# -*- coding: utf-8 -*-
"""
accounting_files_driver.py — Organizador contable 📦 EXPORTACIÓN (spec 13.5).
Archivo NUEVO (Zero-Impact). Lo usa routers/accounting_files.py.

Guarda en Postgres (bytea, D-135-01) los libros que arma el motor 13.4 y los
documentos que sube el usuario: jamás en el bucket público (R-135-05).

  generar(receta, usuario, …)  → motor 13.4 + folio EXP-AAAA-NNNN + SHA-256 + huella
                                 de datos; purga y archiva (snapshot inmutable, R-135-03)
  prevuelo(receta)             → el sello del motor sin archivar (= la carátula, CA-135-02)
  listar / resumen / ficha / descargar / vista_previa / subir / actualizar / eliminar
  calcular_huellas(…)          → vigencia en lote, UNA consulta (§4.6)
  carpetas, tipos documentales, paquetes guardados, matriz de cierres

El submódulo no calcula cifras: las del archivo son las del motor (R-135-06).
Errores: ValueError → 400 · NoEncontrado → 404 · Conflicto → 409 ·
es_tabla_faltante(e) → 503 (falta correr scripts/migrate_exports.py).
"""
import calendar
import csv
import hashlib
import io
import json
import os
import re
import zipfile
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple


def _numero_env(nombre: str, defecto: float) -> float:
    try:
        return float(os.getenv(nombre) or defecto)
    except ValueError:
        return defecto


RETENCION_DIAS = max(1, int(_numero_env("ANALYTICS_EXPORT_RETENCION_DIAS", 90)))
MAX_MB = max(1.0, _numero_env("ANALYTICS_EXPORT_MAX_MB", 10))
MAX_BYTES = int(MAX_MB * 1024 * 1024)
AVISO_VENCE_DIAS = 15            # "⏳ por vencer": le quedan ≤ 15 días
LIMITE_LISTA = 200
LIMITE_VIGENCIA = 500            # tope de archivos que se comparan en una consulta
MAX_COLS_VISTA = 40
MAX_RECETA_BYTES = 200_000       # una relación de 5000 ids cabe holgada

MIME_PDF = "application/pdf"
MIME_PNG = "image/png"
MIME_JPG = "image/jpeg"
MIME_CSV = "text/csv"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

ORIGEN_GENERADO = "GENERADO"
ORIGEN_SUBIDO = "SUBIDO"

# Tipos documentales por defecto (las "categorías" de RRHH, versión contable).
TIPOS_DEFAULT = (
    # clave,        nombre,                icono, color,     orden
    ("libros",      "Libros",              "📘", "#1d4ed8", 1),
    ("estados",     "Estados financieros", "📊", "#0f766e", 2),
    ("impuestos",   "Impuestos",           "🧾", "#b45309", 3),
    ("cartera",     "Cartera",             "💰", "#15803d", 4),
    ("bancos",      "Bancos",              "🏦", "#4338ca", 5),
    ("relaciones",  "Relaciones de TXs",   "📋", "#7c3aed", 6),
    ("soportes",    "Soportes",            "📎", "#64748b", 7),
)

# Paquetes predefinidos (§4.2). Las hojas son claves del motor 13.4; la
# carátula la pone siempre el motor. "completo" = todas las del modo período.
HOJAS_COMPLETO = ["diario", "mayor", "balance_prueba", "estado_resultados", "balance_general",
                  "movimientos", "auxiliar_tercero", "cartera", "impuestos"]
PAQUETES: Dict[str, Dict[str, Any]] = {
    "cierre_mes": {"nombre": "Cierre de mes", "icono": "📘", "tipo": "libros",
                   "hojas": ["diario", "mayor", "balance_prueba", "estado_resultados", "balance_general"]},
    "banco": {"nombre": "Banco / crédito", "icono": "🏦", "tipo": "estados",
              "hojas": ["estado_resultados", "balance_general"]},
    "impuestos": {"nombre": "Impuestos / DIAN", "icono": "🧾", "tipo": "impuestos",
                  "hojas": ["auxiliar_tercero", "impuestos"]},
    "cartera": {"nombre": "Cartera", "icono": "💰", "tipo": "cartera",
                "hojas": ["cartera", "auxiliar_tercero"]},
    "movimientos": {"nombre": "Movimientos", "icono": "📋", "tipo": "libros",
                    "hojas": ["movimientos"]},
    "completo": {"nombre": "Libro completo", "icono": "🧮", "tipo": "libros",
                 "hojas": list(HOJAS_COMPLETO)},
}
PAQUETE_RELACION = "relacion"        # modo transacciones
PAQUETE_LIBRE = "personalizado"      # hojas elegidas a mano
TIPO_RELACION = "relaciones"
TIPO_LIBRE = "libros"
TIPO_SUBIDA = "soportes"

MESES = ("ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic")
_COLOR = re.compile(r"^#[0-9a-fA-F]{6}$")


class NoEncontrado(LookupError):
    """→ 404"""


class Conflicto(Exception):
    """→ 409"""


def es_tabla_faltante(e: BaseException) -> bool:
    """La migración 13.5 no se ha corrido (42P01 = relación inexistente)."""
    return getattr(e, "pgcode", None) == "42P01"


# ══ DDL (la corre scripts/migrate_exports.py; el server NO la ejecuta) ══════

DDL = [
    """
    CREATE TABLE IF NOT EXISTS accounting_doc_types (
        id SERIAL PRIMARY KEY,
        clave VARCHAR(30) UNIQUE,                 -- solo los default; NULL en los propios
        nombre VARCHAR(60) NOT NULL,
        icono VARCHAR(16),
        color VARCHAR(9) NOT NULL DEFAULT '#64748b',
        orden SMALLINT NOT NULL DEFAULT 100,
        es_default BOOLEAN NOT NULL DEFAULT FALSE,
        creado_por VARCHAR(120),
        creado_en TIMESTAMPTZ NOT NULL DEFAULT NOW()
    )
    """,
    "CREATE UNIQUE INDEX IF NOT EXISTS uq_accounting_doc_types_nombre "
    "ON accounting_doc_types (lower(btrim(nombre)))",
    """
    CREATE TABLE IF NOT EXISTS accounting_folders (
        id SERIAL PRIMARY KEY,
        nombre VARCHAR(80) NOT NULL,
        color VARCHAR(9) NOT NULL DEFAULT '#64748b',
        parent_id INTEGER REFERENCES accounting_folders(id) ON DELETE CASCADE,
        portfolio_id INTEGER REFERENCES portfolios(id) ON DELETE SET NULL,
        creado_por VARCHAR(120),
        creado_en TIMESTAMPTZ NOT NULL DEFAULT NOW()
    )
    """,
    "CREATE SEQUENCE IF NOT EXISTS accounting_files_folio_seq",
    """
    CREATE TABLE IF NOT EXISTS accounting_files (
        id SERIAL PRIMARY KEY,
        origen VARCHAR(10) NOT NULL CHECK (origen IN ('GENERADO', 'SUBIDO')),
        folio VARCHAR(20) UNIQUE,                 -- solo generados: EXP-AAAA-NNNN
        nombre VARCHAR(160) NOT NULL,
        nombre_archivo VARCHAR(200) NOT NULL,     -- el de la descarga
        tipo_documental_id INTEGER REFERENCES accounting_doc_types(id) ON DELETE SET NULL,
        paquete VARCHAR(40),
        receta JSONB,                             -- modo transacciones: incluye tx_ids
        portfolio_id INTEGER REFERENCES portfolios(id) ON DELETE SET NULL,
        periodo_desde DATE,
        periodo_hasta DATE,
        folder_id INTEGER REFERENCES accounting_folders(id) ON DELETE SET NULL,
        archivo BYTEA NOT NULL,
        mime_type VARCHAR(100) NOT NULL,
        tamano_bytes INTEGER NOT NULL,
        sha256 CHAR(64) NOT NULL,
        hojas TEXT[],
        sello JSONB,
        huella_datos JSONB,
        creado_por VARCHAR(120),
        creado_en TIMESTAMPTZ NOT NULL DEFAULT NOW(),
        fijado BOOLEAN NOT NULL DEFAULT FALSE,
        nota TEXT,
        descargas INTEGER NOT NULL DEFAULT 0,
        ultima_descarga TIMESTAMPTZ,
        reemplaza_a INTEGER REFERENCES accounting_files(id) ON DELETE SET NULL
    )
    """,
    "CREATE INDEX IF NOT EXISTS idx_accounting_files_empresa_periodo "
    "ON accounting_files (portfolio_id, periodo_hasta)",
    "CREATE INDEX IF NOT EXISTS idx_accounting_files_folder ON accounting_files (folder_id)",
    "CREATE INDEX IF NOT EXISTS idx_accounting_files_tipo ON accounting_files (tipo_documental_id)",
    "CREATE INDEX IF NOT EXISTS idx_accounting_files_purga "
    "ON accounting_files (creado_en) WHERE origen = 'GENERADO' AND NOT fijado",
    """
    CREATE TABLE IF NOT EXISTS analytics_export_paquetes (
        id SERIAL PRIMARY KEY,
        nombre VARCHAR(80) NOT NULL,
        receta JSONB NOT NULL,
        creado_por VARCHAR(120),
        creado_en TIMESTAMPTZ NOT NULL DEFAULT NOW()
    )
    """,
]

SQL_TIPOS_DEFAULT = """
    INSERT INTO accounting_doc_types (clave, nombre, icono, color, orden, es_default, creado_por)
    VALUES (%s, %s, %s, %s, %s, TRUE, 'migracion')
    ON CONFLICT (clave) DO NOTHING
"""


# ══ Conexión ═════════════════════════════════════════════════════════════

@contextmanager
def _conexion(conn=None):
    """Con `conn` dado (tests) lo usa tal cual; si no, presta una del pool y
    hace commit al salir bien o rollback si algo falla."""
    if conn is not None:
        yield conn
        return
    from fin_sys_core.db_pool import get_conn, put_conn
    c = get_conn()
    try:
        yield c
        c.commit()
    except Exception:
        try:
            c.rollback()
        except Exception:
            pass
        raise
    finally:
        put_conn(c)


def _motor():
    from fin_sys_core import export_xlsx
    return export_xlsx


def _filas(cur, columnas) -> List[Dict[str, Any]]:
    return [dict(zip(columnas, r)) for r in cur.fetchall()]


def _fila(cur, columnas) -> Optional[Dict[str, Any]]:
    r = cur.fetchone()
    return dict(zip(columnas, r)) if r else None


def _iso(v):
    if isinstance(v, datetime):
        return v.isoformat(timespec="seconds")
    if isinstance(v, date):
        return v.isoformat()
    return v


def _quien(usuario: Optional[Dict[str, Any]]) -> str:
    u = usuario or {}
    nombre = (u.get("name") or "").strip()
    return (nombre or (f"usuario {u.get('uid')}" if u.get("uid") else "desconocido"))[:120]


def _dec(x) -> str:
    try:
        return str(Decimal(str(x if x is not None else 0)).quantize(Decimal("0.01")))
    except InvalidOperation:
        return "0.00"


def _cop(s) -> str:
    """'1234567.50' → '1.234.567,50' (formato colombiano)."""
    txt = f"{Decimal(_dec(s)):,.2f}"
    return txt.replace(",", "§").replace(".", ",").replace("§", ".")


def _entero(v, campo: str, opcional: bool = True) -> Optional[int]:
    if v in (None, ""):
        if opcional:
            return None
        raise ValueError(f"Falta '{campo}'.")
    try:
        n = int(v)
    except (TypeError, ValueError):
        raise ValueError(f"'{campo}' debe ser un número entero.")
    if n <= 0:
        raise ValueError(f"'{campo}' debe ser positivo.")
    return n


def _texto(v, campo: str, maximo: int, obligatorio: bool = False) -> Optional[str]:
    s = re.sub(r"[\x00-\x1f\x7f]", " ", str(v or "")).strip()
    if not s:
        if obligatorio:
            raise ValueError(f"Falta '{campo}'.")
        return None
    if len(s) > maximo:
        raise ValueError(f"'{campo}' admite máximo {maximo} caracteres.")
    return s


def _color(v, defecto: str = "#64748b") -> str:
    s = str(v or "").strip()
    if not s:
        return defecto
    if not _COLOR.match(s):
        raise ValueError("El color debe ser hexadecimal, p. ej. #1d4ed8.")
    return s.lower()


# ══ Paquetes y receta ════════════════════════════════════════════════════

def paquetes_publicos() -> List[Dict[str, Any]]:
    return [{"clave": k, "nombre": v["nombre"], "icono": v["icono"], "tipo": v["tipo"],
             "hojas": list(v["hojas"])} for k, v in PAQUETES.items()]


def preparar_receta(receta: Optional[Dict[str, Any]], paquete: Optional[str]) -> Tuple[Dict[str, Any], str]:
    """Paquete → hojas (si no vinieron). Devuelve (receta para el motor, clave)."""
    r = dict(receta or {})
    modo = str(r.get("modo") or "periodo").strip().lower()
    if modo == "transacciones":
        return r, PAQUETE_RELACION
    clave = str(paquete or "").strip().lower() or PAQUETE_LIBRE
    if clave not in PAQUETES and clave != PAQUETE_LIBRE:
        raise ValueError(f"Paquete desconocido: {clave!r}. Usa {', '.join(PAQUETES)} o '{PAQUETE_LIBRE}'.")
    if clave in PAQUETES and not r.get("hojas"):
        r["hojas"] = list(PAQUETES[clave]["hojas"])
    return r, clave


def paquete_efectivo(clave: str, r: Dict[str, Any]) -> str:
    """Si las hojas normalizadas no son exactamente las del paquete, el archivo
    es 'personalizado' (la matriz de cierres solo cuenta paquetes completos)."""
    if clave not in PAQUETES:
        return clave
    propias = set(PAQUETES[clave]["hojas"])
    elegidas = set(r.get("hojas") or []) - {"caratula"}
    return clave if elegidas == propias else PAQUETE_LIBRE


def receta_json(r: Dict[str, Any], relativo: Optional[str] = None) -> Dict[str, Any]:
    """Receta normalizada → JSON guardable (fechas ISO; sin folio)."""
    out = {k: _iso(v) for k, v in r.items() if k != "folio"}
    out["filtros"] = dict(r.get("filtros") or {})
    if relativo:
        out["relativo"] = str(relativo)[:30]
    return out


def alcance_de_receta(receta: Dict[str, Any]) -> Dict[str, Any]:
    """Qué datos cubre un archivo, para su huella (§4.6). Los filtros finos no
    cuentan: la huella vigila los libros del período completo."""
    if (receta.get("modo") or "periodo") == "transacciones":
        return {"tx_ids": sorted({int(i) for i in receta.get("tx_ids") or []})}
    pids = sorted({int(p) for p in receta.get("portfolios") or []})
    return {"pids": pids or None, "desde": _iso(receta.get("desde")), "hasta": _iso(receta.get("hasta"))}


def etiqueta_periodo(desde: Optional[date], hasta: Optional[date]) -> str:
    if hasta is None:
        return "sin período"
    ultimo = calendar.monthrange(hasta.year, hasta.month)[1]
    fin_de_mes = hasta.day == ultimo
    if desde and fin_de_mes and desde.day == 1:
        if (desde.year, desde.month) == (hasta.year, hasta.month):
            return f"{MESES[hasta.month - 1]} {hasta.year}"
        if desde.year == hasta.year and desde.month == 1 and hasta.month == 12:
            return f"año {hasta.year}"
        if (desde.year == hasta.year and desde.month % 3 == 1
                and hasta.month == desde.month + 2):
            return f"T{(hasta.month + 2) // 3} {hasta.year}"
    fin = hasta.strftime("%d/%m/%Y")
    return f"{desde.strftime('%d/%m/%Y')} – {fin}" if desde else f"hasta {fin}"


def nombre_por_defecto(clave: str, r: Dict[str, Any], datos: Dict[str, Any]) -> str:
    """Nombre visible cuando el usuario no dio uno: 'Cierre de mes · Pegasus · sep 2026'."""
    if r["modo"] == "transacciones":
        return f"Relación de {len(datos.get('txs') or [])} transacciones"
    base = PAQUETES.get(clave, {}).get("nombre") or "Libro personalizado"
    empresas = ", ".join(e["nombre"] for e in datos.get("empresas") or []) or "Consolidado"
    return f"{base} · {empresas} · {etiqueta_periodo(r.get('desde'), r.get('hasta'))}"[:160]


def _fecha_dmy(s) -> Optional[date]:
    try:
        return datetime.strptime(str(s), "%d/%m/%Y").date()
    except (TypeError, ValueError):
        return None


# ══ Huella de datos y vigencia (§4.6) ════════════════════════════════════

_SQL_HUELLAS = """
SELECT s.id,
       tx.n, tx.suma, tx.max_id, tx.ids,
       je.n, je.debitos, je.max_id
  FROM jsonb_to_recordset(%s::jsonb)
       AS s(id int, pids int[], desde date, hasta date, tx_ids int[])
 CROSS JOIN LATERAL (
       SELECT COUNT(*) AS n, COALESCE(SUM(t.net_value), 0) AS suma, MAX(t.id) AS max_id,
              CASE WHEN s.tx_ids IS NULL THEN NULL
                   ELSE COALESCE(array_agg(t.id ORDER BY t.id), '{}'::int[]) END AS ids
         FROM transactions t
        WHERE CASE WHEN s.tx_ids IS NOT NULL THEN t.id = ANY(s.tx_ids)
                   ELSE (COALESCE(cardinality(s.pids), 0) = 0 OR t.portfolio_id = ANY(s.pids))
                        AND (s.desde IS NULL OR t.transaction_date >= s.desde)
                        AND (s.hasta IS NULL OR t.transaction_date <= s.hasta) END
 ) tx
 CROSS JOIN LATERAL (
       SELECT COUNT(*) AS n, COALESCE(SUM(j.debito), 0) AS debitos, MAX(j.id) AS max_id
         FROM kernel_journal_entries j
        WHERE j.estado IN ('CONTABILIZADO', 'ANULADO')
          AND CASE WHEN s.tx_ids IS NOT NULL THEN j.tx_id = ANY(s.tx_ids)
                   ELSE (COALESCE(cardinality(s.pids), 0) = 0 OR j.portfolio_id = ANY(s.pids))
                        AND (s.desde IS NULL OR j.fecha >= s.desde)
                        AND (s.hasta IS NULL OR j.fecha <= s.hasta) END
 ) je
"""


def calcular_huellas(cur, alcances: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Huella de N alcances en UNA consulta (tarjetas visibles, matriz de cierres)."""
    if not alcances:
        return []
    payload = [{"id": i, "pids": a.get("pids"), "desde": a.get("desde"), "hasta": a.get("hasta"),
                "tx_ids": a.get("tx_ids")} for i, a in enumerate(alcances)]
    cur.execute(_SQL_HUELLAS, (json.dumps(payload),))
    por: Dict[int, Dict[str, Any]] = {}
    for i, n_tx, suma, max_tx, ids, n_ln, debitos, max_ln in cur.fetchall():
        h = {"n_txs": int(n_tx or 0), "suma_neto": _dec(suma), "max_tx": max_tx,
             "n_lineas": int(n_ln or 0), "debitos": _dec(debitos), "max_linea": max_ln}
        if alcances[i].get("tx_ids") is not None:
            h["tx_vivas"] = sorted(int(x) for x in (ids or []))
        por[int(i)] = h
    return [por.get(i, {}) for i in range(len(alcances))]


def comparar_huella(antes: Optional[Dict[str, Any]], ahora: Dict[str, Any]) -> Tuple[str, List[str]]:
    """→ ('VIGENTE' | 'CAMBIO' | 'SIN_HUELLA', detalles en español)."""
    if not antes:
        return "SIN_HUELLA", []
    d: List[str] = []
    if "tx_vivas" in antes:                      # relación de transacciones
        perdidas = sorted(set(antes.get("tx_vivas") or []) - set(ahora.get("tx_vivas") or []))
        for i in perdidas[:5]:
            d.append(f"la TX #{i} se eliminó")
        if len(perdidas) > 5:
            d.append(f"y {len(perdidas) - 5} TX(s) eliminadas más")
        if not perdidas and antes.get("suma_neto") != ahora.get("suma_neto"):
            d.append(f"cambió el valor neto de las TXs: {_cop(antes.get('suma_neto'))} → "
                     f"{_cop(ahora.get('suma_neto'))}")
    else:                                        # período
        dn = int(ahora.get("n_txs", 0)) - int(antes.get("n_txs", 0))
        if dn > 0:
            d.append(f"+{dn} TX(s) registrada(s) después")
        elif dn < 0:
            d.append(f"{-dn} TX(s) eliminada(s) del período")
        elif ahora.get("max_tx") != antes.get("max_tx"):
            d.append("hay TXs nuevas y otras eliminadas en el período")
        elif antes.get("suma_neto") != ahora.get("suma_neto"):
            d.append(f"cambió el valor neto del período: {_cop(antes.get('suma_neto'))} → "
                     f"{_cop(ahora.get('suma_neto'))}")
    clave_asientos = ("n_lineas", "debitos", "max_linea")
    if any(antes.get(k) != ahora.get(k) for k in clave_asientos):
        dl = int(ahora.get("n_lineas", 0)) - int(antes.get("n_lineas", 0))
        if dl:
            d.append(f"los asientos en libros cambiaron ({'+' if dl > 0 else ''}{dl} líneas)")
        else:
            d.append("los asientos en libros cambiaron (mismas líneas, otros valores)")
    return ("CAMBIO" if d else "VIGENTE"), d


def _vigencias(cur, filas: List[Dict[str, Any]]) -> Dict[int, Dict[str, Any]]:
    gen = [f for f in filas if f.get("origen") == ORIGEN_GENERADO and f.get("huella_datos")]
    gen = gen[:LIMITE_VIGENCIA]
    ahora = calcular_huellas(cur, [alcance_de_receta(f.get("receta") or {}) for f in gen])
    out: Dict[int, Dict[str, Any]] = {}
    for f, h in zip(gen, ahora):
        estado, detalles = comparar_huella(f["huella_datos"], h)
        out[f["id"]] = {"estado": estado, "detalles": detalles}
    return out


# ══ Folio y purga ════════════════════════════════════════════════════════

def siguiente_folio(cur, anio: int) -> str:
    """Secuencia global: único y creciente aunque se purguen archivos. Puede
    tener huecos si una generación falla tras pedirlo (como toda secuencia)."""
    cur.execute("SELECT nextval('accounting_files_folio_seq')")
    return f"EXP-{anio}-{int(cur.fetchone()[0]):04d}"


SQL_PURGA = """
    DELETE FROM accounting_files
     WHERE origen = 'GENERADO' AND NOT fijado
       AND creado_en < NOW() - make_interval(days => %s)
"""


def purgar(cur) -> int:
    """D-135-04: lo GENERADO no fijado vive RETENCION_DIAS; fijados y subidos, siempre."""
    cur.execute(SQL_PURGA, (RETENCION_DIAS,))
    return max(int(cur.rowcount or 0), 0)


def _id_tipo(cur, clave: str) -> Optional[int]:
    cur.execute("SELECT id FROM accounting_doc_types WHERE clave = %s", (clave,))
    r = cur.fetchone()
    return int(r[0]) if r else None


def _validar_destino(cur, folder_id: Optional[int], tipo_id: Optional[int],
                     portfolio_id: Optional[int] = None) -> None:
    for tabla, valor, que in (("accounting_folders", folder_id, "La carpeta"),
                              ("accounting_doc_types", tipo_id, "El tipo documental"),
                              ("portfolios", portfolio_id, "La empresa")):
        if valor is None:
            continue
        cur.execute(f"SELECT 1 FROM {tabla} WHERE id = %s", (valor,))
        if not cur.fetchone():
            raise ValueError(f"{que} {valor} no existe.")


# ══ Generar y pre-vuelo ══════════════════════════════════════════════════

def prevuelo(receta: Optional[Dict[str, Any]], paquete: Optional[str] = None,
             usuario: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """El mismo camino que generar, sin folio ni archivo: lo que dirá la carátula."""
    motor = _motor()
    r_in, clave = preparar_receta(receta, paquete)
    r = motor.normalizar_receta(r_in)
    datos = motor.recolectar(r)
    contenido, sello = motor.construir_libro(datos, r, usuario=usuario)
    return {"sello": sello, "paquete": paquete_efectivo(clave, r),
            "tamano_bytes": len(contenido), "excede_tope": len(contenido) > MAX_BYTES,
            "max_mb": MAX_MB}


def generar(receta: Optional[Dict[str, Any]], usuario: Optional[Dict[str, Any]] = None, *,
            paquete: Optional[str] = None, folder_id: Optional[int] = None,
            tipo_documental_id: Optional[int] = None, reemplaza_a: Optional[int] = None,
            conn=None) -> Dict[str, Any]:
    motor = _motor()
    r_in, clave = preparar_receta(receta, paquete)
    r = motor.normalizar_receta(r_in)
    clave = paquete_efectivo(clave, r)
    guardable = receta_json(r, (receta or {}).get("relativo"))
    folder_id = _entero(folder_id, "folder_id")
    tipo_documental_id = _entero(tipo_documental_id, "tipo_documental_id")

    # 1) Huella ANTES de leer los datos: si algo entra en medio, el archivo
    #    marcará ⚠ (falsa alarma) en vez de ✅ (falsa tranquilidad).
    with _conexion(conn) as c:
        cur = c.cursor()
        _validar_destino(cur, folder_id, tipo_documental_id)
        huella = calcular_huellas(cur, [alcance_de_receta(guardable)])[0]
        pids_tx: List[int] = []
        if r["modo"] == "transacciones":
            cur.execute("SELECT DISTINCT portfolio_id FROM transactions WHERE id = ANY(%s)", (r["tx_ids"],))
            pids_tx = sorted(int(x[0]) for x in cur.fetchall() if x[0] is not None)
        cur.close()

    # 2) Datos (el motor usa sus propias conexiones).
    datos = motor.recolectar(r)

    # 3) Folio → libro → archivo, en una transacción. El nombre de descarga
    #    sale del nombre que dio el usuario; el visible (carátula) puede ser
    #    el derivado del paquete, la empresa y el período.
    with _conexion(conn) as c:
        cur = c.cursor()
        r["folio"] = siguiente_folio(cur, motor.hoy_colombia().year)
        nombre_archivo = f"{r['folio']}_{motor.nombre_archivo(r, datos)}"[:200]
        nombre = r.get("nombre") or nombre_por_defecto(clave, r, datos)
        r["nombre"] = nombre[:120]
        contenido, sello = motor.construir_libro(datos, r, usuario=usuario)
        if len(contenido) > MAX_BYTES:
            raise ValueError(f"El libro pesa {len(contenido) / 1048576:.1f} MB y el tope es {MAX_MB:g} MB: "
                             "pide un período más corto o menos hojas.")
        pids = r["portfolios"] or pids_tx
        if r["modo"] == "transacciones":
            desde = _fecha_dmy(sello.get("rango_real", {}).get("desde"))
            hasta = _fecha_dmy(sello.get("rango_real", {}).get("hasta"))
        else:
            desde, hasta = r["desde"], r["hasta"]
        tipo = tipo_documental_id or _id_tipo(
            cur, TIPO_RELACION if clave == PAQUETE_RELACION else PAQUETES.get(clave, {}).get("tipo", TIPO_LIBRE))
        sha =hashlib.sha256(contenido).hexdigest()
        purgados = purgar(cur)
        cur.execute("""
            INSERT INTO accounting_files
                (origen, folio, nombre, nombre_archivo, tipo_documental_id, paquete, receta,
                 portfolio_id, periodo_desde, periodo_hasta, folder_id, archivo, mime_type,
                 tamano_bytes, sha256, hojas, sello, huella_datos, creado_por, reemplaza_a)
            VALUES ('GENERADO', %s, %s, %s, %s, %s, %s::jsonb, %s, %s, %s, %s, %s, %s,
                    %s, %s, %s, %s::jsonb, %s::jsonb, %s, %s)
            RETURNING id, creado_en
        """, (r["folio"], nombre, nombre_archivo, tipo, clave,
              json.dumps(guardable, ensure_ascii=False, default=str),
              pids[0] if len(pids) == 1 else None, desde, hasta, folder_id, contenido, MIME_XLSX,
              len(contenido), sha, list(sello.get("hojas") or []),
              json.dumps(sello, ensure_ascii=False, default=str),
              json.dumps(huella), _quien(usuario), reemplaza_a))
        nuevo_id, creado_en = cur.fetchone()
        cur.close()
    return {"id": int(nuevo_id), "folio": r["folio"], "nombre": nombre, "nombre_archivo": nombre_archivo,
            "paquete": clave, "tamano_bytes": len(contenido), "sha256": sha, "sello": sello,
            "creado_en": _iso(creado_en), "purgados": purgados}


def regenerar(file_id: int, usuario: Optional[Dict[str, Any]] = None, desde: Optional[str] = None,
              hasta: Optional[str] = None, conn=None) -> Dict[str, Any]:
    """Misma receta con datos de hoy (o período nuevo) → folio NUEVO con
    reemplaza_a; el original queda intacto (CA-135-05)."""
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute("""SELECT origen, receta, paquete, folder_id, tipo_documental_id
                         FROM accounting_files WHERE id = %s""", (file_id,))
        f = _fila(cur, ("origen", "receta", "paquete", "folder_id", "tipo_documental_id"))
        cur.close()
    if not f:
        raise NoEncontrado(f"El archivo {file_id} no existe.")
    if f["origen"] != ORIGEN_GENERADO or not f["receta"]:
        raise ValueError("Solo se regeneran los archivos que FIN-SYS generó.")
    receta = dict(f["receta"])
    if desde or hasta:
        if receta.get("modo") == "transacciones":
            raise ValueError("Una relación de transacciones se regenera con sus mismas TXs, sin período.")
        receta["desde"], receta["hasta"] = desde or None, hasta or None
        receta.pop("nombre", None)      # período nuevo → nombre nuevo
    paquete = f["paquete"] if f["paquete"] in PAQUETES else None
    return generar(receta, usuario, paquete=paquete, folder_id=f["folder_id"],
                   tipo_documental_id=f["tipo_documental_id"], reemplaza_a=file_id, conn=conn)


# ══ Subir documentos externos ════════════════════════════════════════════

def _es_xlsx(contenido: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            return "xl/workbook.xml" in z.namelist()
    except (zipfile.BadZipFile, OSError):
        return False


def detectar_tipo(contenido: bytes, nombre_archivo: str) -> Tuple[str, str]:
    """El tipo se decide por el CONTENIDO (firma), no por lo que diga el cliente."""
    if not contenido:
        raise ValueError("El archivo está vacío.")
    if len(contenido) > MAX_BYTES:
        raise ValueError(f"El archivo pesa {len(contenido) / 1048576:.1f} MB; el tope es {MAX_MB:g} MB.")
    ext = os.path.splitext(nombre_archivo or "")[1].lower().lstrip(".")
    if contenido.startswith(b"%PDF-"):
        return MIME_PDF, "pdf"
    if contenido.startswith(b"\x89PNG\r\n\x1a\n"):
        return MIME_PNG, "png"
    if contenido.startswith(b"\xff\xd8\xff"):
        return MIME_JPG, "jpg"
    if contenido.startswith(b"PK\x03\x04"):
        if ext == "xlsx" and _es_xlsx(contenido):
            return MIME_XLSX, "xlsx"
        raise ValueError("De los comprimidos solo se acepta el libro de Excel .xlsx.")
    if ext == "csv" and b"\x00" not in contenido[:65536]:
        return MIME_CSV, "csv"
    raise ValueError("Tipo de archivo no permitido: sube PDF, PNG, JPG, XLSX o CSV.")


def _nombre_seguro(nombre: str, ext: str) -> str:
    base = os.path.basename(str(nombre or "").replace("\\", "/"))
    base = os.path.splitext(base)[0]
    base = re.sub(r'[\x00-\x1f\x7f"<>:/\\|?*]+', "_", base).strip(" ._") or "documento"
    return f"{base[:180]}.{ext}"


def _periodo_mes(anio: Optional[int], mes: Optional[int]) -> Tuple[Optional[date], Optional[date]]:
    if anio is None:
        if mes is not None:
            raise ValueError("Si eliges el mes, elige también el año.")
        return None, None
    if not 2000 <= anio <= 2100:
        raise ValueError("Año fuera de rango.")
    if mes is None:
        return date(anio, 1, 1), date(anio, 12, 31)
    if not 1 <= mes <= 12:
        raise ValueError("El mes va de 1 a 12.")
    return date(anio, mes, 1), date(anio, mes, calendar.monthrange(anio, mes)[1])


def subir(contenido: bytes, nombre_original: str, meta: Optional[Dict[str, Any]] = None,
          usuario: Optional[Dict[str, Any]] = None, conn=None) -> Dict[str, Any]:
    m = meta or {}
    mime, ext = detectar_tipo(contenido, nombre_original)
    nombre_archivo = _nombre_seguro(nombre_original, ext)
    nombre = _texto(m.get("nombre"), "nombre", 160) or os.path.splitext(nombre_archivo)[0][:160]
    nota = _texto(m.get("nota"), "nota", 2000)
    folder_id = _entero(m.get("folder_id"), "folder_id")
    tipo_id = _entero(m.get("tipo_documental_id"), "tipo_documental_id")
    portfolio_id = _entero(m.get("portfolio_id"), "portfolio_id")
    desde, hasta = _periodo_mes(_entero(m.get("anio"), "anio"), _entero(m.get("mes"), "mes"))
    sha = hashlib.sha256(contenido).hexdigest()
    with _conexion(conn) as c:
        cur = c.cursor()
        _validar_destino(cur, folder_id, tipo_id, portfolio_id)
        tipo_id = tipo_id or _id_tipo(cur, TIPO_SUBIDA)
        purgados = purgar(cur)
        cur.execute("""
            INSERT INTO accounting_files
                (origen, nombre, nombre_archivo, tipo_documental_id, portfolio_id, periodo_desde,
                 periodo_hasta, folder_id, archivo, mime_type, tamano_bytes, sha256, creado_por, nota)
            VALUES ('SUBIDO', %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            RETURNING id, creado_en
        """, (nombre, nombre_archivo, tipo_id, portfolio_id, desde, hasta, folder_id, contenido,
              mime, len(contenido), sha, _quien(usuario), nota))
        nuevo_id, creado_en = cur.fetchone()
        cur.close()
    return {"id": int(nuevo_id), "nombre": nombre, "nombre_archivo": nombre_archivo, "mime_type": mime,
            "tamano_bytes": len(contenido), "sha256": sha, "creado_en": _iso(creado_en),
            "purgados": purgados}


# ══ Listar, resumen, ficha ═══════════════════════════════════════════════

_COLS = ("id", "origen", "folio", "nombre", "nombre_archivo", "tipo_documental_id", "paquete",
         "portfolio_id", "empresa", "portafolio", "periodo_desde", "periodo_hasta", "folder_id",
         "mime_type", "tamano_bytes", "sha256", "hojas", "creado_por", "creado_en", "fijado", "nota",
         "descargas", "ultima_descarga", "reemplaza_a", "receta", "huella_datos")
def _select(extra: str = "") -> str:
    return f"""
    SELECT f.id, f.origen, f.folio, f.nombre, f.nombre_archivo, f.tipo_documental_id, f.paquete,
           f.portfolio_id,
           (SELECT e.name FROM entities e WHERE e.portfolio_id = f.portfolio_id ORDER BY e.id LIMIT 1),
           p.name, f.periodo_desde, f.periodo_hasta, f.folder_id, f.mime_type, f.tamano_bytes,
           f.sha256, f.hojas, f.creado_por, f.creado_en, f.fijado, f.nota, f.descargas,
           f.ultima_descarga, f.reemplaza_a, f.receta, f.huella_datos{extra}
      FROM accounting_files f
      LEFT JOIN portfolios p ON p.id = f.portfolio_id
"""


_SELECT = _select()
_POR_VENCER = ("f.origen = 'GENERADO' AND NOT f.fijado "
               "AND f.creado_en < NOW() - make_interval(days => %s)")


def _publico(f: Dict[str, Any], vig: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    out = {k: _iso(v) for k, v in f.items() if k not in ("receta", "huella_datos")}
    out["empresa"] = (f.get("empresa") or f.get("portafolio") or None)
    out["modo"] = (f.get("receta") or {}).get("modo")
    out["hojas"] = list(f.get("hojas") or [])
    vence = None
    if f.get("origen") == ORIGEN_GENERADO and not f.get("fijado") and isinstance(f.get("creado_en"), datetime):
        vence = (f["creado_en"] + timedelta(days=RETENCION_DIAS)).date().isoformat()
    out["vence_el"] = vence
    out["vigencia"] = vig
    return out


def _where(filtros: Dict[str, Any]) -> Tuple[str, List[Any]]:
    cond, params = [], []
    if "folder_id" in filtros and filtros["folder_id"] is not None:
        cond.append("f.folder_id = %s")
        params.append(int(filtros["folder_id"]))
    if filtros.get("portfolio_id") is not None:
        if int(filtros["portfolio_id"]) == 0:          # 0 = consolidado / sin empresa
            cond.append("f.portfolio_id IS NULL")
        else:
            cond.append("f.portfolio_id = %s")
            params.append(int(filtros["portfolio_id"]))
    if filtros.get("anio") is not None:
        cond.append("EXTRACT(YEAR FROM f.periodo_hasta) = %s")
        params.append(int(filtros["anio"]))
    if filtros.get("mes") is not None:
        cond.append("EXTRACT(MONTH FROM f.periodo_hasta) = %s")
        params.append(int(filtros["mes"]))
    if filtros.get("tipo") is not None:
        if int(filtros["tipo"]) == 0:                  # 0 = sin tipo documental
            cond.append("f.tipo_documental_id IS NULL")
        else:
            cond.append("f.tipo_documental_id = %s")
            params.append(int(filtros["tipo"]))
    if filtros.get("origen"):
        o = str(filtros["origen"]).upper()
        if o not in (ORIGEN_GENERADO, ORIGEN_SUBIDO):
            raise ValueError("'origen' es GENERADO o SUBIDO.")
        cond.append("f.origen = %s")
        params.append(o)
    if filtros.get("fijado"):
        cond.append("f.fijado")
    if filtros.get("por_vencer"):
        cond.append(_POR_VENCER)
        params.append(max(RETENCION_DIAS - AVISO_VENCE_DIAS, 0))
    q = (filtros.get("q") or "").strip()
    if q:
        cond.append("(f.nombre ILIKE %s OR f.folio ILIKE %s OR f.nombre_archivo ILIKE %s "
                    "OR COALESCE(f.nota, '') ILIKE %s)")
        patron = "%" + q.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"
        params += [patron] * 4
    return (" WHERE " + " AND ".join(cond)) if cond else "", params


def listar(filtros: Optional[Dict[str, Any]] = None, conn=None) -> Dict[str, Any]:
    """Lista del organizador (sin bytes) + vigencia en lote de lo listado."""
    filtros = dict(filtros or {})
    limite = max(1, min(int(filtros.get("limit") or 50), LIMITE_LISTA))
    offset = max(0, int(filtros.get("offset") or 0))
    where, params = _where(filtros)
    orden = " ORDER BY f.creado_en DESC, f.id DESC"
    with _conexion(conn) as c:
        cur = c.cursor()
        if filtros.get("cambiaron"):
            cur.execute(_SELECT + where + (" AND " if where else " WHERE ") + "f.origen = 'GENERADO'"
                        + orden + " LIMIT %s", params + [LIMITE_VIGENCIA])
            filas = _filas(cur, _COLS)
            vig = _vigencias(cur, filas)
            filas = [f for f in filas if vig.get(f["id"], {}).get("estado") == "CAMBIO"]
            total = len(filas)
            filas = filas[offset:offset + limite]
        else:
            cur.execute("SELECT COUNT(*) FROM accounting_files f" + where, params)
            total = int(cur.fetchone()[0])
            cur.execute(_SELECT + where + orden + " LIMIT %s OFFSET %s", params + [limite, offset])
            filas = _filas(cur, _COLS)
            vig = _vigencias(cur, filas)
        cur.close()
    return {"items": [_publico(f, vig.get(f["id"])) for f in filas], "total": total,
            "limit": limite, "offset": offset}


def resumen(conn=None) -> Dict[str, Any]:
    """Cabecera del desplegable + conteos del lateral + árbol automático
    Empresa → Año → Mes ("12 archivos · 3,1 MB · ⚠ 2 cambiaron")."""
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute(f"""
            SELECT COUNT(*), COALESCE(SUM(f.tamano_bytes), 0),
                   COUNT(*) FILTER (WHERE f.fijado),
                   COUNT(*) FILTER (WHERE {_POR_VENCER}),
                   COUNT(*) FILTER (WHERE f.origen = 'SUBIDO')
              FROM accounting_files f
        """, (max(RETENCION_DIAS - AVISO_VENCE_DIAS, 0),))
        n, peso, fijados, por_vencer, subidos = cur.fetchone()
        cur.execute("SELECT tipo_documental_id, COUNT(*) FROM accounting_files GROUP BY 1")
        por_tipo = {str(t if t is not None else 0): int(k) for t, k in cur.fetchall()}
        cur.execute("SELECT folder_id, COUNT(*) FROM accounting_files WHERE folder_id IS NOT NULL GROUP BY 1")
        por_carpeta = {str(fid): int(k) for fid, k in cur.fetchall()}
        cur.execute("""
            SELECT f.portfolio_id,
                   (SELECT e.name FROM entities e WHERE e.portfolio_id = f.portfolio_id ORDER BY e.id LIMIT 1),
                   p.name,
                   EXTRACT(YEAR FROM f.periodo_hasta)::int, EXTRACT(MONTH FROM f.periodo_hasta)::int,
                   COUNT(*)
              FROM accounting_files f
              LEFT JOIN portfolios p ON p.id = f.portfolio_id
             GROUP BY 1, 2, 3, 4, 5
             ORDER BY 3 NULLS FIRST, 4 DESC NULLS LAST, 5 DESC NULLS LAST
        """)
        arbol = [{"portfolio_id": pid, "empresa": (emp or port or ("Consolidado" if pid is None else f"Empresa {pid}")),
                  "anio": anio, "mes": mes, "n": int(k)}
                 for pid, emp, port, anio, mes, k in cur.fetchall()]
        cur.execute("SELECT id, origen, receta, huella_datos FROM accounting_files "
                    "WHERE origen = 'GENERADO' ORDER BY creado_en DESC LIMIT %s", (LIMITE_VIGENCIA,))
        vig = _vigencias(cur, _filas(cur, ("id", "origen", "receta", "huella_datos")))
        cur.close()
    return {"archivos": int(n), "bytes": int(peso), "fijados": int(fijados), "por_vencer": int(por_vencer),
            "subidos": int(subidos), "cambiaron": sum(1 for v in vig.values() if v["estado"] == "CAMBIO"),
            "por_tipo": por_tipo, "por_carpeta": por_carpeta, "arbol": arbol,
            "retencion_dias": RETENCION_DIAS, "max_mb": MAX_MB}


def ficha(file_id: int, conn=None) -> Dict[str, Any]:
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute(_select(", f.sello") + " WHERE f.id = %s", (file_id,))
        f = _fila(cur, _COLS + ("sello",))
        if not f:
            cur.close()
            raise NoEncontrado(f"El archivo {file_id} no existe.")
        vig = _vigencias(cur, [f]).get(f["id"])
        cur.execute("SELECT id, folio, creado_en FROM accounting_files WHERE reemplaza_a = %s ORDER BY id",
                    (file_id,))
        sucesores = [{"id": i, "folio": fo, "creado_en": _iso(ce)} for i, fo, ce in cur.fetchall()]
        anterior = None
        if f.get("reemplaza_a"):
            cur.execute("SELECT id, folio FROM accounting_files WHERE id = %s", (f["reemplaza_a"],))
            r = cur.fetchone()
            anterior = {"id": r[0], "folio": r[1]} if r else None
        cur.close()
    out = _publico(f, vig)
    out["sello"] = f.get("sello")
    out["receta"] = f.get("receta")
    out["reemplazado_por"] = sucesores
    out["reemplaza"] = anterior
    return out


# ══ Descargar y vista previa ═════════════════════════════════════════════

def descargar(file_id: int, conn=None) -> Dict[str, Any]:
    """El archivo exacto (mismo SHA-256 que al generar) y cuenta la descarga."""
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute("""
            UPDATE accounting_files SET descargas = descargas + 1, ultima_descarga = NOW()
             WHERE id = %s
         RETURNING archivo, mime_type, nombre_archivo, sha256
        """, (file_id,))
        r = cur.fetchone()
        cur.close()
    if not r:
        raise NoEncontrado(f"El archivo {file_id} no existe.")
    return {"contenido": bytes(r[0]), "mime_type": r[1], "nombre_archivo": r[2], "sha256": r[3]}


def _celda(v, formula=None):
    if v is None:
        if isinstance(formula, str) and formula.startswith("="):
            return formula
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(timespec="minutes")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, Decimal):
        return float(v)
    return str(v)


def _recortar(filas: List[List[Any]]) -> Tuple[List[List[Any]], int]:
    ancho = 0
    for f in filas:
        for i in range(len(f) - 1, -1, -1):
            if f[i] not in ("", None):
                ancho = max(ancho, i + 1)
                break
    ancho = min(ancho, MAX_COLS_VISTA)
    return [(list(f) + [""] * ancho)[:ancho] for f in filas], ancho


def vista_xlsx(contenido: bytes, hoja: Optional[str] = None, limite: int = 50) -> Dict[str, Any]:
    """Primeras filas de una hoja. Las celdas con fórmula que Excel no ha
    calculado (las de FIN-SYS) muestran la fórmula: '=SUM(D5:D9)'."""
    from openpyxl import load_workbook
    limite = max(1, min(int(limite or 50), 500))
    wb_v = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    wb_f = load_workbook(io.BytesIO(contenido), read_only=True, data_only=False)
    try:
        nombres = list(wb_v.sheetnames)
        if not nombres:
            return {"tipo": "tabla", "hojas": [], "hoja": None, "columnas": 0, "filas": [], "mas_filas": False}
        elegida = hoja if hoja in nombres else nombres[0]
        filas, mas = [], False
        pares = zip(wb_v[elegida].iter_rows(max_row=limite + 1, values_only=True),
                    wb_f[elegida].iter_rows(max_row=limite + 1, values_only=True))
        for i, (fv, ff) in enumerate(pares):
            if i >= limite:
                mas = True
                break
            ff = list(ff) + [None] * max(0, len(fv) - len(ff))
            filas.append([_celda(v, f) for v, f in zip(fv, ff)][:MAX_COLS_VISTA])
        filas, ancho = _recortar(filas)
        return {"tipo": "tabla", "hojas": nombres, "hoja": elegida, "columnas": ancho,
                "filas": filas, "mas_filas": mas}
    finally:
        wb_v.close()
        wb_f.close()


def vista_csv(contenido: bytes, limite: int = 50) -> Dict[str, Any]:
    limite = max(1, min(int(limite or 50), 500))
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = contenido.decode("latin-1")
    muestra = texto[:8192]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
    filas, mas = [], False
    for i, fila in enumerate(csv.reader(io.StringIO(texto), dialecto)):
        if i >= limite:
            mas = True
            break
        filas.append(fila[:MAX_COLS_VISTA])
    filas, ancho = _recortar(filas)
    return {"tipo": "tabla", "hojas": [], "hoja": None, "columnas": ancho, "filas": filas, "mas_filas": mas}


def vista_previa(file_id: int, hoja: Optional[str] = None, limite: int = 50, conn=None) -> Dict[str, Any]:
    """xlsx/csv → {'tipo': 'tabla', …}; PDF/imagen → {'tipo': 'binario', contenido, mime_type}."""
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute("SELECT archivo, mime_type, nombre_archivo FROM accounting_files WHERE id = %s", (file_id,))
        r = cur.fetchone()
        cur.close()
    if not r:
        raise NoEncontrado(f"El archivo {file_id} no existe.")
    contenido, mime, nombre = bytes(r[0]), r[1], r[2]
    if mime == MIME_XLSX:
        return vista_xlsx(contenido, hoja, limite)
    if mime == MIME_CSV:
        return vista_csv(contenido, limite)
    return {"tipo": "binario", "contenido": contenido, "mime_type": mime, "nombre_archivo": nombre}


# ══ Editar y eliminar ════════════════════════════════════════════════════

def actualizar(file_id: int, cambios: Dict[str, Any], conn=None) -> Dict[str, Any]:
    """PATCH: nombre, nota, fijado, carpeta (null = sacarlo), tipo documental."""
    sets, params = [], []
    if "nombre" in cambios:
        sets.append("nombre = %s")
        params.append(_texto(cambios["nombre"], "nombre", 160, obligatorio=True))
    if "nota" in cambios:
        sets.append("nota = %s")
        params.append(_texto(cambios["nota"], "nota", 2000))
    if "fijado" in cambios:
        if not isinstance(cambios["fijado"], bool):
            raise ValueError("'fijado' es verdadero o falso.")
        sets.append("fijado = %s")
        params.append(cambios["fijado"])
    folder_id = tipo_id = None
    if "folder_id" in cambios:
        folder_id = _entero(cambios["folder_id"], "folder_id")
        sets.append("folder_id = %s")
        params.append(folder_id)
    if "tipo_documental_id" in cambios:
        tipo_id = _entero(cambios["tipo_documental_id"], "tipo_documental_id")
        sets.append("tipo_documental_id = %s")
        params.append(tipo_id)
    if not sets:
        raise ValueError("Nada que cambiar: usa nombre, nota, fijado, folder_id o tipo_documental_id.")
    with _conexion(conn) as c:
        cur = c.cursor()
        _validar_destino(cur, folder_id, tipo_id)
        cur.execute(f"UPDATE accounting_files SET {', '.join(sets)} WHERE id = %s RETURNING id",
                    params + [file_id])
        ok = cur.fetchone()
        cur.close()
    if not ok:
        raise NoEncontrado(f"El archivo {file_id} no existe.")
    return ficha(file_id, conn=conn)


def eliminar(file_id: int, conn=None) -> Dict[str, Any]:
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute("DELETE FROM accounting_files WHERE id = %s RETURNING id, folio, nombre", (file_id,))
        r = cur.fetchone()
        cur.close()
    if not r:
        raise NoEncontrado(f"El archivo {file_id} no existe.")
    return {"eliminado": True, "id": r[0], "folio": r[1], "nombre": r[2]}


# ══ Carpetas propias ═════════════════════════════════════════════════════

_COLS_CARPETA = ("id", "nombre", "color", "parent_id", "portfolio_id", "creado_por", "creado_en", "n_archivos")


def listar_carpetas(conn=None) -> List[Dict[str, Any]]:
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute("""
            SELECT c.id, c.nombre, c.color, c.parent_id, c.portfolio_id, c.creado_por, c.creado_en,
                   (SELECT COUNT(*) FROM accounting_files f WHERE f.folder_id = c.id)
              FROM accounting_folders c
             ORDER BY lower(c.nombre), c.id
        """)
        filas = _filas(cur, _COLS_CARPETA)
        cur.close()
    return [{k: _iso(v) for k, v in f.items()} for f in filas]


def crear_carpeta(datos: Dict[str, Any], usuario: Optional[Dict[str, Any]] = None, conn=None) -> Dict[str, Any]:
    nombre = _texto(datos.get("nombre"), "nombre", 80, obligatorio=True)
    color = _color(datos.get("color"))
    parent_id = _entero(datos.get("parent_id"), "parent_id")
    portfolio_id = _entero(datos.get("portfolio_id"), "portfolio_id")
    with _conexion(conn) as c:
        cur = c.cursor()
        _validar_destino(cur, parent_id, None, portfolio_id)
        cur.execute("""
            INSERT INTO accounting_folders (nombre, color, parent_id, portfolio_id, creado_por)
            VALUES (%s, %s, %s, %s, %s)
            RETURNING id, nombre, color, parent_id, portfolio_id, creado_por, creado_en, 0
        """, (nombre, color, parent_id, portfolio_id, _quien(usuario)))
        f = _fila(cur, _COLS_CARPETA)
        cur.close()
    return {k: _iso(v) for k, v in f.items()}


def actualizar_carpeta(folder_id: int, datos: Dict[str, Any], conn=None) -> Dict[str, Any]:
    sets, params = [], []
    if "nombre" in datos:
        sets.append("nombre = %s")
        params.append(_texto(datos["nombre"], "nombre", 80, obligatorio=True))
    if "color" in datos:
        sets.append("color = %s")
        params.append(_color(datos["color"]))
    nuevo_padre = None
    if "parent_id" in datos:
        nuevo_padre = _entero(datos["parent_id"], "parent_id")
        if nuevo_padre == folder_id:
            raise ValueError("Una carpeta no puede quedar dentro de sí misma.")
        sets.append("parent_id = %s")
        params.append(nuevo_padre)
    if not sets:
        raise ValueError("Nada que cambiar: usa nombre, color o parent_id.")
    with _conexion(conn) as c:
        cur = c.cursor()
        if nuevo_padre is not None:
            _validar_destino(cur, nuevo_padre, None)
            cur.execute("""
                WITH RECURSIVE ancestros AS (
                    SELECT id, parent_id FROM accounting_folders WHERE id = %s
                    UNION ALL
                    SELECT a.id, a.parent_id FROM accounting_folders a
                      JOIN ancestros x ON a.id = x.parent_id
                )
                SELECT 1 FROM ancestros WHERE id = %s LIMIT 1
            """, (nuevo_padre, folder_id))
            if cur.fetchone():
                raise ValueError("No puedes mover una carpeta dentro de una de sus subcarpetas.")
        cur.execute(f"""
            UPDATE accounting_folders SET {', '.join(sets)} WHERE id = %s
         RETURNING id, nombre, color, parent_id, portfolio_id, creado_por, creado_en,
                   (SELECT COUNT(*) FROM accounting_files f WHERE f.folder_id = accounting_folders.id)
        """, params + [folder_id])
        f = _fila(cur, _COLS_CARPETA)
        cur.close()
    if not f:
        raise NoEncontrado(f"La carpeta {folder_id} no existe.")
    return {k: _iso(v) for k, v in f.items()}


def eliminar_carpeta(folder_id: int, conn=None) -> Dict[str, Any]:
    """Borra la carpeta y sus subcarpetas; los archivos NO se borran: quedan
    sueltos en su ubicación automática (empresa → año → mes) y por tipo."""
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute("DELETE FROM accounting_folders WHERE id = %s RETURNING id, nombre", (folder_id,))
        r = cur.fetchone()
        cur.close()
    if not r:
        raise NoEncontrado(f"La carpeta {folder_id} no existe.")
    return {"eliminada": True, "id": r[0], "nombre": r[1]}


# ══ Tipos documentales ═══════════════════════════════════════════════════

_COLS_TIPO = ("id", "clave", "nombre", "icono", "color", "orden", "es_default", "n_archivos")


def listar_tipos(conn=None) -> List[Dict[str, Any]]:
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute("""
            SELECT t.id, t.clave, t.nombre, t.icono, t.color, t.orden, t.es_default,
                   (SELECT COUNT(*) FROM accounting_files f WHERE f.tipo_documental_id = t.id)
              FROM accounting_doc_types t
             ORDER BY t.orden, lower(t.nombre), t.id
        """)
        filas = _filas(cur, _COLS_TIPO)
        cur.close()
    return filas


def _nombre_repetido(e: BaseException) -> bool:
    return getattr(e, "pgcode", None) == "23505"


def crear_tipo(datos: Dict[str, Any], usuario: Optional[Dict[str, Any]] = None, conn=None) -> Dict[str, Any]:
    nombre = _texto(datos.get("nombre"), "nombre", 60, obligatorio=True)
    icono = _texto(datos.get("icono"), "icono", 16)
    color = _color(datos.get("color"))
    try:
        with _conexion(conn) as c:
            cur = c.cursor()
            cur.execute("""
                INSERT INTO accounting_doc_types (nombre, icono, color, orden, es_default, creado_por)
                VALUES (%s, %s, %s, (SELECT COALESCE(MAX(orden), 0) + 1 FROM accounting_doc_types), FALSE, %s)
                RETURNING id, clave, nombre, icono, color, orden, es_default, 0
            """, (nombre, icono, color, _quien(usuario)))
            f = _fila(cur, _COLS_TIPO)
            cur.close()
    except Exception as e:
        if _nombre_repetido(e):
            raise Conflicto(f"Ya existe un tipo documental llamado «{nombre}».")
        raise
    return f


def actualizar_tipo(tipo_id: int, datos: Dict[str, Any], conn=None) -> Dict[str, Any]:
    sets, params = [], []
    if "nombre" in datos:
        sets.append("nombre = %s")
        params.append(_texto(datos["nombre"], "nombre", 60, obligatorio=True))
    if "icono" in datos:
        sets.append("icono = %s")
        params.append(_texto(datos["icono"], "icono", 16))
    if "color" in datos:
        sets.append("color = %s")
        params.append(_color(datos["color"]))
    if "orden" in datos:
        orden = _entero(datos["orden"], "orden", opcional=False)
        sets.append("orden = %s")
        params.append(min(orden, 32000))
    if not sets:
        raise ValueError("Nada que cambiar: usa nombre, icono, color u orden.")
    try:
        with _conexion(conn) as c:
            cur = c.cursor()
            cur.execute(f"""
                UPDATE accounting_doc_types SET {', '.join(sets)} WHERE id = %s
             RETURNING id, clave, nombre, icono, color, orden, es_default,
                       (SELECT COUNT(*) FROM accounting_files f WHERE f.tipo_documental_id = accounting_doc_types.id)
            """, params + [tipo_id])
            f = _fila(cur, _COLS_TIPO)
            cur.close()
    except Exception as e:
        if _nombre_repetido(e):
            raise Conflicto("Ya existe un tipo documental con ese nombre.")
        raise
    if not f:
        raise NoEncontrado(f"El tipo documental {tipo_id} no existe.")
    return f


def eliminar_tipo(tipo_id: int, conn=None) -> Dict[str, Any]:
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute("SELECT es_default, nombre FROM accounting_doc_types WHERE id = %s", (tipo_id,))
        r = cur.fetchone()
        if not r:
            cur.close()
            raise NoEncontrado(f"El tipo documental {tipo_id} no existe.")
        if r[0]:
            cur.close()
            raise ValueError(f"«{r[1]}» es un tipo por defecto: se puede renombrar, no borrar.")
        cur.execute("DELETE FROM accounting_doc_types WHERE id = %s", (tipo_id,))
        cur.close()
    return {"eliminado": True, "id": tipo_id, "nombre": r[1]}


# ══ Paquetes guardados (recetas propias) ═════════════════════════════════

def listar_paquetes(conn=None) -> Dict[str, Any]:
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute("SELECT id, nombre, receta, creado_por, creado_en FROM analytics_export_paquetes "
                    "ORDER BY lower(nombre), id")
        guardados = [{k: _iso(v) for k, v in f.items()}
                     for f in _filas(cur, ("id", "nombre", "receta", "creado_por", "creado_en"))]
        cur.close()
    return {"predefinidos": paquetes_publicos(), "guardados": guardados}


def guardar_paquete(nombre: str, receta: Dict[str, Any], usuario: Optional[Dict[str, Any]] = None,
                    conn=None) -> Dict[str, Any]:
    nombre = _texto(nombre, "nombre", 80, obligatorio=True)
    if not isinstance(receta, dict) or not receta:
        raise ValueError("La receta debe ser un objeto con modo, empresas, período y hojas.")
    modo = str(receta.get("modo") or "periodo").lower()
    if modo not in ("periodo", "transacciones"):
        raise ValueError("Modo desconocido en la receta.")
    texto = json.dumps(receta, ensure_ascii=False, default=str)
    if len(texto.encode()) > MAX_RECETA_BYTES:
        raise ValueError("La receta es demasiado grande.")
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute("""
            INSERT INTO analytics_export_paquetes (nombre, receta, creado_por)
            VALUES (%s, %s::jsonb, %s) RETURNING id, nombre, receta, creado_por, creado_en
        """, (nombre, texto, _quien(usuario)))
        f = _fila(cur, ("id", "nombre", "receta", "creado_por", "creado_en"))
        cur.close()
    return {k: _iso(v) for k, v in f.items()}


def eliminar_paquete(paquete_id: int, conn=None) -> Dict[str, Any]:
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute("DELETE FROM analytics_export_paquetes WHERE id = %s RETURNING id, nombre", (paquete_id,))
        r = cur.fetchone()
        cur.close()
    if not r:
        raise NoEncontrado(f"El paquete {paquete_id} no existe.")
    return {"eliminado": True, "id": r[0], "nombre": r[1]}


# ══ Matriz de cierres (vista 🗓) ═════════════════════════════════════════

def cierres(portfolio_id: Optional[int], anio: int, conn=None) -> Dict[str, Any]:
    """Último folio por paquete × mes de UNA empresa (o consolidado si None).
    Cuenta solo exportaciones de un mes exacto con un paquete completo."""
    if not 2000 <= int(anio) <= 2100:
        raise ValueError("Año fuera de rango.")
    with _conexion(conn) as c:
        cur = c.cursor()
        cur.execute("""
            SELECT DISTINCT ON (f.paquete, EXTRACT(MONTH FROM f.periodo_hasta))
                   f.id, f.origen, f.folio, f.paquete, EXTRACT(MONTH FROM f.periodo_hasta)::int,
                   f.creado_en, f.fijado, f.receta, f.huella_datos
              FROM accounting_files f
             WHERE f.origen = 'GENERADO' AND f.paquete = ANY(%s)
               AND f.portfolio_id IS NOT DISTINCT FROM %s
               AND EXTRACT(YEAR FROM f.periodo_hasta) = %s
               AND f.periodo_desde = date_trunc('month', f.periodo_hasta)::date
               AND f.periodo_hasta = (date_trunc('month', f.periodo_hasta) + interval '1 month - 1 day')::date
             ORDER BY f.paquete, EXTRACT(MONTH FROM f.periodo_hasta), f.creado_en DESC, f.id DESC
        """, (list(PAQUETES), portfolio_id, int(anio)))
        filas = _filas(cur, ("id", "origen", "folio", "paquete", "mes", "creado_en", "fijado",
                             "receta", "huella_datos"))
        vig = _vigencias(cur, filas)
        cur.close()
    celdas: Dict[str, Dict[str, Any]] = {k: {} for k in PAQUETES}
    for f in filas:
        v = vig.get(f["id"]) or {"estado": "SIN_HUELLA", "detalles": []}
        celdas[f["paquete"]][str(f["mes"])] = {"id": f["id"], "folio": f["folio"], "fijado": f["fijado"],
                                              "creado_en": _iso(f["creado_en"]), **v}
    return {"anio": int(anio), "portfolio_id": portfolio_id, "paquetes": paquetes_publicos(), "celdas": celdas}
