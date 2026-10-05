# -*- coding: utf-8 -*-
"""Integración del motor de libros .xlsx contra Supabase (CA-134-01/02/05).
SOLO LECTURA: arma libros reales y los cruza con el kernel; no escribe nada.

  · el archivo abre (openpyxl lo relee sin error);
  · el diario del libro = balance de prueba del kernel (Σ débitos y créditos);
  · el Mayor cierra igual que el balance (sin advertencia de descuadre);
  · consolidado ⊇ empresas: la suma de TXs por empresa = TXs del consolidado.

Ejecutar:  .venv\\Scripts\\python.exe -m unittest tests.test_export_xlsx_db -v
"""
import io
import os
import sys
import unittest

from openpyxl import load_workbook

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "fin_sys_core"))

_env = os.path.join(_ROOT, ".env")
if os.path.exists(_env):
    with open(_env, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, _, v = line.partition("=")
                os.environ.setdefault(k.strip(), v.strip())

import fin_sys_core  # noqa: E402,F401
from fin_sys_core import export_xlsx as ex  # noqa: E402

DESDE = "2026-01-01"


def _db_disponible():
    try:
        from db_pool import get_conn, put_conn
        conn = get_conn()
        try:
            cur = conn.cursor()
            cur.execute("SELECT 1 FROM kernel_journal_entries LIMIT 1")
            cur.close()
            return True
        finally:
            put_conn(conn)
    except Exception:
        return False


_DB_OK = _db_disponible()


@unittest.skipUnless(_DB_OK, "BD o kernel_journal_entries no disponibles")
class TestLibroReal(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        from kernel import kernel_reports as kr
        cls.contenido, cls.sello, cls.archivo = ex.armar_libro({"desde": DESDE}, usuario={"name": "test-db"})
        cls.kernel_bp = kr.balance_prueba(None, DESDE, cls.sello["hasta"])

    def test_el_archivo_abre_y_trae_las_diez_hojas(self):  # CA-134-01
        wb = load_workbook(io.BytesIO(self.contenido))
        self.assertEqual(wb.sheetnames, list(ex.HOJAS_PERIODO.values()))
        self.assertTrue(self.archivo.startswith("FINSYS_CONSOLIDADO_2026-01-01_"))

    def test_diario_del_libro_igual_al_balance_del_kernel(self):  # CA-134-02
        tc = self.sello["totales_control"]
        tot = self.kernel_bp["totales"]
        self.assertAlmostEqual(tc["debitos"], tot["mov_debito"], places=2)
        self.assertAlmostEqual(tc["creditos"], tot["mov_credito"], places=2)
        self.assertEqual(tc["cuadra"], self.kernel_bp["cuadra"])
        adv = " | ".join(self.sello["advertencias"])
        self.assertNotIn("no coincide con el balance de prueba", adv)
        self.assertNotIn("Mayor no coincide", adv)

    def test_consolidado_contiene_a_cada_empresa(self):  # CA-134-05
        from database_driver import obtener_portafolios
        suma = 0
        for p in obtener_portafolios():
            _c, s, _a = ex.armar_libro({"desde": DESDE, "portfolio_id": p["id"],
                                        "hojas": "diario"}, usuario={"name": "test-db"})
            suma += s["n_txs"]
            self.assertLessEqual(s["totales_control"]["debitos"],
                                 self.sello["totales_control"]["debitos"] + 0.01)
        self.assertEqual(suma, self.sello["n_txs"])


if __name__ == "__main__":
    unittest.main()
