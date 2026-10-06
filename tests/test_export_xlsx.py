# -*- coding: utf-8 -*-
"""Motor de libros .xlsx (spec 13.4), SIN BD.

Arma libros con datos de prueba, los RELEE con openpyxl y verifica lo que un
contador revisaría: que las hojas estén, que los totales sean fórmulas =SUM
reales (se evalúan aquí con un mini evaluador), que el diario y el balance
cuadren, que el USD nunca se sume al COP, que los filtros no toquen los libros
oficiales y que la carátula diga la verdad (sello y advertencias).
También cubre la recolección con las fuentes parcheadas y el endpoint con
TestClient (guards reales: 401 sin sesión, 403 sin rol contador/admin).

Ejecutar:  .venv\\Scripts\\python.exe -m unittest tests.test_export_xlsx -v
"""
import io
import os
import re
import sys
import unittest
from datetime import date, datetime
from unittest import mock

from openpyxl import load_workbook

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "fin_sys_core"))

import fin_sys_core  # noqa: E402,F401  (un solo objeto por módulo)
from fin_sys_core import export_xlsx as ex  # noqa: E402

AHORA = datetime(2026, 10, 5, 16, 10)
URL = "https://x.supabase.co/storage/v1/object/public/hr-docs/evidence/a.jpg"


# ── Mini evaluador de las fórmulas que escribe el motor ─────────────────

_REF = re.compile(r"\b([A-Z]{1,3})(\d+)\b")
_SUM = re.compile(r"SUM\(([A-Z]{1,3})(\d+):([A-Z]{1,3})(\d+)\)")
_IF = re.compile(r'^IF\(ROUND\((.+),2\)=0,"(.*)","(.*)"\)$')


def valor(ws, coord):
    c = ws[coord]
    if c.data_type == "f" and isinstance(c.value, str):
        return evaluar(ws, c.value[1:])
    return c.value if c.value is not None else 0


def evaluar(ws, expr):
    m = _IF.match(expr)
    if m:
        return m.group(2) if round(evaluar(ws, m.group(1)), 2) == 0 else m.group(3)

    def suma(mm):
        col, a, b = mm.group(1), int(mm.group(2)), int(mm.group(4))
        return repr(sum(float(valor(ws, f"{col}{r}") or 0) for r in range(a, b + 1)))

    expr = _SUM.sub(suma, expr)
    expr = _REF.sub(lambda mm: repr(float(valor(ws, mm.group(0)) or 0)), expr)
    return eval(expr, {"__builtins__": {}})   # solo números y + − tras sustituir


def buscar(ws, texto, col=None):
    """Primera celda cuyo valor es exactamente `texto` → (fila, columna)."""
    for fila in ws.iter_rows():
        for c in fila:
            if c.value == texto and (col is None or c.column == col):
                return c.row, c.column
    raise AssertionError(f"No encontré {texto!r} en {ws.title}")


def textos(ws):
    return [c.value for fila in ws.iter_rows() for c in fila if isinstance(c.value, str)]


# ── Datos de prueba (lo que devolvería recolectar) ──────────────────────

def _tx(i, tipo, neto, fecha, **kw):
    base = {"id": i, "type": tipo, "amount": neto, "net_value": neto, "concept": f"TX {i}",
            "transaction_date": fecha, "category": "General", "tax_iva_amount": 0, "tax_gmf_amount": 0,
            "account_id": 1, "account_name": "Bancolombia", "third_party_id": None,
            "third_party_name": None, "transaction_currency": "COP", "trm": 1,
            "portfolio_name": "Negocio A", "evidences": [], "evidence_file_path": None, "tags": None}
    base.update(kw)
    return base


def _grupo(gid, estado, tx_id, fecha, lineas, tercero=None):
    return {"entry_group_id": gid, "estado": estado, "tx_id": tx_id, "fecha": fecha,
            "portfolio_name": "Negocio A", "referencia": f"TX-{tx_id}", "descripcion": f"Asiento TX {tx_id}",
            "tx": {"tercero": tercero} if tercero else None,
            "lineas": [{"id": n, "linea": n, "cuenta_codigo": c, "cuenta_nombre": nom, "debito": d, "credito": cr}
                       for n, (c, nom, d, cr) in enumerate(lineas, start=1)]}


def datos_periodo():
    txs = [
        _tx(1, "INGRESO", 1_000_000, date(2026, 9, 5), category="Ventas", third_party_id=7,
            third_party_name="Sandra Jiménez", identification_type="CC", identification_number="52123",
            tax_iva_amount=190_000, evidences=[URL]),
        _tx(2, "GASTO", 200_000, date(2026, 9, 10), category="Viajes", account_id=None,
            account_name=None, concept='=HYPERLINK("http://malo")', tax_iva_amount=38_000,
            tax_gmf_amount=800, third_party_id=8, third_party_name="Hotel Dann"),
        _tx(3, "GASTO", 50, date(2026, 9, 12), category="Software", transaction_currency="USD", trm=4000),
    ]
    asientos = [
        _grupo("JE-1", "CONTABILIZADO", 1, "2026-09-05",
               [("1110", "Bancos", 1_000_000, 0), ("4135", "Ventas", 0, 1_000_000)], tercero="Sandra Jiménez"),
        _grupo("JE-2", "CONTABILIZADO", 2, "2026-09-10",
               [("5105", "Gastos de viaje", 200_000, 0), ("1110", "Bancos", 0, 200_000)]),
        _grupo("JE-3", "BORRADOR", 3, "2026-09-12",
               [("5195", "Software", 999, 0), ("1110", "Bancos", 0, 999)]),
    ]
    bp = [
        {"codigo": "1110", "nombre": "Bancos", "tipo": "ACTIVO", "ini": 500_000.0, "db": 1_000_000.0, "cr": 200_000.0},
        {"codigo": "3105", "nombre": "Capital", "tipo": "PATRIMONIO", "ini": -500_000.0, "db": 0.0, "cr": 0.0},
        {"codigo": "4135", "nombre": "Ventas", "tipo": "INGRESO", "ini": 0.0, "db": 0.0, "cr": 1_000_000.0},
        {"codigo": "5105", "nombre": "Gastos de viaje", "tipo": "GASTO", "ini": 0.0, "db": 200_000.0, "cr": 0.0},
    ]
    return {
        "empresas": [{"id": 1, "portafolio": "Negocio A", "nombre": "Finanzas Julian", "nit": "900.123.456-7"}],
        "alias": {"Negocio A": "Finanzas Julian"}, "advertencias_bd": [],
        "txs": txs, "asientos": asientos, "truncado": False, "bp": bp,
        "er": {"ingresos": [{"codigo": "4135", "nombre": "Ventas", "tipo": "INGRESO", "saldo": 1_000_000.0}],
               "gastos": [{"codigo": "5105", "nombre": "Gastos de viaje", "tipo": "GASTO", "saldo": 200_000.0}]},
        "bg": {"activos": [{"codigo": "1110", "nombre": "Bancos", "tipo": "ACTIVO", "saldo": 1_300_000.0}],
               "pasivos": [],
               "patrimonio": [{"codigo": "3105", "nombre": "Capital", "tipo": "PATRIMONIO", "saldo": 500_000.0}],
               "utilidad": 800_000.0},
        "cartera": [
            {"type": "CXC", "third_party_id": 7, "third_party_name": "Sandra Jiménez", "identification_number": "52123",
             "concept": "Factura 12", "due_date": date(2026, 8, 20), "original_amount": 300_000,
             "remaining_balance": 100_000, "transaction_date": date(2026, 8, 1)},
            {"type": "CXP", "third_party_id": 8, "third_party_name": "Hotel Dann", "identification_number": "",
             "concept": "Reserva", "due_date": date(2026, 10, 15), "original_amount": 50_000,
             "remaining_balance": 50_000, "transaction_date": date(2026, 9, 1)},
            {"type": "CXC", "third_party_id": 7, "third_party_name": "Sandra Jiménez", "concept": "Pagada",
             "due_date": date(2026, 7, 1), "original_amount": 10_000, "remaining_balance": 0},
        ],
    }


def receta(**kw):
    base = {"portfolio_id": 1, "desde": "2026-09-01", "hasta": "2026-09-30"}
    base.update(kw)
    return ex.normalizar_receta(base)


def libro(datos=None, r=None, usuario=None):
    contenido, sello = ex.construir_libro(datos or datos_periodo(), r or receta(),
                                          usuario=usuario or {"name": "Andrés", "role": "owner"}, ahora=AHORA)
    return load_workbook(io.BytesIO(contenido)), sello


# ══ Receta ═══════════════════════════════════════════════════════════════

class TestReceta(unittest.TestCase):

    def test_defaults_del_modo_periodo(self):
        r = ex.normalizar_receta({})
        self.assertEqual(r["modo"], "periodo")
        self.assertEqual(r["hojas"], list(ex.HOJAS_PERIODO))
        self.assertEqual(r["portfolios"], [])
        self.assertIsNone(r["desde"])
        self.assertEqual(r["hasta"], ex.hoy_colombia())

    def test_empresa_unica_y_lista_se_unen_sin_duplicar(self):
        self.assertEqual(ex.normalizar_receta({"portfolio_id": 2, "portfolios": "2, 3"})["portfolios"], [2, 3])

    def test_subconjunto_de_hojas_lleva_caratula_y_respeta_el_orden(self):
        r = ex.normalizar_receta({"hojas": "impuestos,DIARIO"})
        self.assertEqual(r["hojas"], ["caratula", "diario", "impuestos"])

    def test_errores_del_usuario_son_valueerror(self):
        malas = [
            {"hojas": "diario,inventadas"},
            {"desde": "2026-10-01", "hasta": "2026-09-01"},
            {"desde": "ayer"},
            {"filtros": {"tipos": "PRESTAMO"}},
            {"nivel_puc": "detalle"},
            {"modo": "transacciones"},
            {"tx_ids": "1,2"},
            {"modo": "otro"},
            {"portfolios": "uno"},
            {"modo": "transacciones", "tx_ids": list(range(1, ex.MAX_TX_IDS + 2))},
        ]
        for m in malas:
            with self.subTest(m=m), self.assertRaises(ValueError):
                ex.normalizar_receta(m)

    def test_modo_transacciones_no_inventa_fechas(self):
        r = ex.normalizar_receta({"modo": "transacciones", "tx_ids": "5 6 5"})
        self.assertEqual(r["tx_ids"], [5, 6])
        self.assertIsNone(r["hasta"])
        self.assertEqual(r["hojas"], list(ex.HOJAS_TRANSACCIONES))


# ══ Libro por período ═════════════════════════════════════════════════════

class TestLibroPeriodo(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.wb, cls.sello = libro()

    def test_hojas_en_el_orden_del_contador(self):
        self.assertEqual(self.wb.sheetnames, list(ex.HOJAS_PERIODO.values()))

    def test_diario_solo_en_libros_y_cuadra_con_formulas(self):  # CA-134-02, CA-134-06
        ws = self.wb["LIBRO DIARIO"]
        fila, _ = buscar(ws, "TOTAL", col=1)
        deb = ws.cell(row=fila, column=8)   # Fecha, Asiento, Estado, Referencia, Descripción, Código, Cuenta, Débito
        self.assertTrue(str(deb.value).startswith("=SUM("))
        self.assertEqual(evaluar(ws, deb.value[1:]), 1_200_000)
        self.assertEqual(evaluar(ws, ws.cell(row=fila, column=9).value[1:]), 1_200_000)
        self.assertEqual(valor(ws, ws.cell(row=fila + 1, column=8).coordinate), "✔ Débitos = Créditos")
        self.assertNotIn("JE-3", textos(ws))           # el BORRADOR no está en los libros
        self.assertEqual(self.sello["n_lineas"], 4)
        self.assertEqual(self.sello["n_asientos"], 2)

    def test_formato_numerico_contable(self):
        ws = self.wb["LIBRO DIARIO"]
        self.assertEqual(ws["H5"].number_format, ex.FMT_NUM)
        self.assertEqual(ws["A5"].number_format, ex.FMT_FECHA)
        self.assertEqual(ws.freeze_panes, "A5")

    def test_balance_de_prueba_cuadra(self):
        ws = self.wb["BALANCE DE PRUEBA"]
        fila, _ = buscar(ws, "TOTAL", col=1)
        self.assertEqual(evaluar(ws, ws[f"F{fila}"].value[1:]), evaluar(ws, ws[f"G{fila}"].value[1:]))
        self.assertEqual(valor(ws, f"F{fila + 1}"), "✔ Movimientos cuadran")
        self.assertEqual(valor(ws, f"F{fila + 2}"), "✔ Saldos cuadran")
        self.assertTrue(self.sello["totales_control"]["cuadra"])

    def test_mayor_encadena_saldos_y_cierra_igual_que_el_balance(self):
        ws = self.wb["MAYOR"]
        fila, _ = buscar(ws, "1110 · Bancos")
        cierre, _ = buscar(ws, "Totales y saldo final")
        self.assertGreater(cierre, fila)
        self.assertEqual(valor(ws, f"F{cierre}"), 1_300_000)   # 500k anterior + 1M − 200k
        self.assertFalse(any("Mayor no coincide" in a for a in self.sello["advertencias"]))

    def test_estado_de_resultados_y_balance_general(self):
        er = self.wb["ESTADO DE RESULTADOS"]
        fila, _ = buscar(er, "UTILIDAD (PÉRDIDA) DEL PERÍODO")
        self.assertEqual(valor(er, f"C{fila}"), 800_000)
        bg = self.wb["BALANCE GENERAL"]
        fila, _ = buscar(bg, "VERIFICACIÓN")
        self.assertEqual(valor(bg, f"B{fila}"), "✔ Activo = Pasivo + Patrimonio")
        tc = self.sello["totales_control"]
        self.assertEqual(tc["utilidad_periodo"], 800_000)
        self.assertTrue(tc["ecuacion_contable"])

    def test_usd_en_bloque_propio_y_nunca_sumado_al_cop(self):  # CA-134-04
        ws = self.wb["MOVIMIENTOS"]
        fila_cop, _ = buscar(ws, "TOTAL COP")
        letras = {ws.cell(row=5, column=c).value: ws.cell(row=5, column=c).column_letter
                  for c in range(1, ws.max_column + 1)}
        self.assertEqual(valor(ws, f"{letras['Ingreso (neto)']}{fila_cop}"), 1_000_000)
        self.assertEqual(valor(ws, f"{letras['Gasto (neto)']}{fila_cop}"), 200_000)   # los 50 USD no entran
        self.assertIn("MOVIMIENTOS EN USD — NO se suman al COP", textos(ws))
        self.assertEqual(self.sello["otras_monedas"], {"USD": 1})

    def test_texto_que_empieza_por_igual_jamas_es_formula(self):
        ws = self.wb["MOVIMIENTOS"]
        fila, col = buscar(ws, '=HYPERLINK("http://malo")')
        self.assertEqual(ws.cell(row=fila, column=col).data_type, "s")

    def test_caratula_con_sello_nit_enlaces_y_advertencias(self):  # CA-134-07
        ws = self.wb["CARÁTULA"]
        t = textos(ws)
        self.assertIn("Finanzas Julian (portafolio interno: Negocio A)", t)
        self.assertIn("900.123.456-7", t)
        self.assertIn("2026-10-05 16:10 (hora Colombia)", t)
        self.assertIn("Andrés (owner)", t)
        fila, _ = buscar(ws, "LIBRO DIARIO", col=1)
        self.assertEqual(ws.cell(row=fila, column=1).hyperlink.location, "'LIBRO DIARIO'!A1")
        adv = " | ".join(self.sello["advertencias"])
        for esperado in ("sin cuenta bancaria", "en USD", "BORRADOR", "sin asiento contabilizado", "saldos pendientes ACTUALES"):
            self.assertIn(esperado, adv)
        self.assertEqual(self.sello["n_txs"], 3)
        self.assertEqual(self.sello["rango_real"], {"desde": "05/09/2026", "hasta": "12/09/2026"})

    def test_cartera_por_edades(self):
        ws = self.wb["CARTERA POR EDADES"]
        f_cxc, _ = buscar(ws, "TOTAL CXC")
        f_cxp, _ = buscar(ws, "TOTAL CXP")
        self.assertEqual(valor(ws, f"G{f_cxc}"), 100_000)   # la pagada (saldo 0) no aparece
        self.assertEqual(valor(ws, f"J{f_cxc}"), 100_000)   # 41 días → 31–60
        self.assertEqual(valor(ws, f"H{f_cxp}"), 50_000)    # vence después del corte → corriente

    def test_impuestos_mes_a_mes(self):
        ws = self.wb["IMPUESTOS"]
        fila, _ = buscar(ws, "TOTAL", col=1)
        self.assertEqual(valor(ws, f"B{fila}"), 190_000)
        self.assertEqual(valor(ws, f"C{fila}"), 38_000)
        self.assertEqual(valor(ws, f"D{fila}"), 152_000)
        self.assertEqual(valor(ws, f"E{fila}"), 800)

    def test_auxiliar_por_tercero(self):
        ws = self.wb["AUXILIAR POR TERCERO"]
        fila, _ = buscar(ws, "Sandra Jiménez")
        self.assertEqual(ws[f"B{fila}"].value, "CC 52123")
        self.assertEqual(valor(ws, f"F{fila}"), 1_000_000)
        total, _ = buscar(ws, "TOTAL", col=1)
        self.assertEqual(valor(ws, f"C{total}"), 2)          # solo COP: la de USD no se cuenta aquí


class TestFiltrosYNivel(unittest.TestCase):

    def test_filtros_finos_no_tocan_los_libros_oficiales(self):  # D-134-07
        wb, sello = libro(r=receta(filtros={"categorias": ["ventas"]}))
        mov = wb["MOVIMIENTOS"]
        fila, _ = buscar(mov, "TOTAL COP")
        letras = {mov.cell(row=4, column=c).value: mov.cell(row=4, column=c).column_letter
                  for c in range(1, mov.max_column + 1)}
        self.assertEqual(valor(mov, f"{letras['Ingreso (neto)']}{fila}"), 1_000_000)
        self.assertEqual(valor(mov, f"{letras['Gasto (neto)']}{fila}"), 0)
        self.assertEqual(sello["n_lineas"], 4)                  # el diario sigue completo
        self.assertEqual((sello["n_txs"], sello["n_txs_filtradas"]), (3, 1))
        self.assertIn("Los libros oficiales", " ".join(textos(wb["CARÁTULA"])))

    def test_nivel_clase_agrupa_el_balance_de_prueba(self):
        wb, sello = libro(r=receta(nivel_puc="clase"))
        ws = wb["BALANCE DE PRUEBA"]
        codigos = [ws[f"A{r}"].value for r in range(5, 9)]
        self.assertEqual(codigos, ["1", "3", "4", "5"])
        self.assertEqual(ws["B5"].value, "Activo")
        fila, _ = buscar(ws, "TOTAL", col=1)
        self.assertEqual(valor(ws, f"F{fila + 2}"), "✔ Saldos cuadran")

    def test_nombres_oficiales_del_plan_de_cuentas(self):  # D-134-09
        d = datos_periodo()
        d["nombres_cuenta"] = {"1110": "Bancos Nacionales"}
        wb, _ = libro(datos=d)
        self.assertIn("1110 · Bancos Nacionales", textos(wb["MAYOR"]))
        self.assertIn("Bancos Nacionales", textos(wb["LIBRO DIARIO"]))
        self.assertIn("Bancos Nacionales", textos(wb["BALANCE GENERAL"]))
        self.assertIn("Ventas", textos(wb["ESTADO DE RESULTADOS"]))   # sin nombre oficial: queda el del asiento

    def test_nivel_grupo_se_bautiza_con_el_puc(self):
        d = datos_periodo()
        d["nombres_cuenta"] = {"11": "DISPONIBLE", "41": "OPERACIONALES"}
        wb, _ = libro(datos=d, r=receta(nivel_puc="grupo"))
        ws = wb["BALANCE DE PRUEBA"]
        self.assertEqual((ws["A5"].value, ws["B5"].value), ("11", "DISPONIBLE"))
        self.assertIn("Grupo 31", textos(ws))                  # sin nombre conocido: rótulo honesto

    def test_mayor_filtrado_por_cuentas_puc(self):
        wb, _ = libro(r=receta(filtros={"cuentas_puc": "41"}))
        t = textos(wb["MAYOR"])
        self.assertIn("4135 · Ventas", t)
        self.assertNotIn("1110 · Bancos", t)


class TestLibroVacioYRelacion(unittest.TestCase):

    def test_libro_vacio_sale_y_lo_dice(self):  # CA-134-08
        vacio = {"empresas": [], "alias": {}, "txs": [], "asientos": [], "bp": [],
                 "er": {"ingresos": [], "gastos": []},
                 "bg": {"activos": [], "pasivos": [], "patrimonio": [], "utilidad": 0}, "cartera": []}
        wb, sello = libro(datos=vacio, r=receta(portfolio_id=None))
        self.assertEqual(len(wb.sheetnames), len(ex.HOJAS_PERIODO))     # todas (12 desde el 06-oct)
        self.assertTrue(any("sale en cero" in a for a in sello["advertencias"]))
        self.assertIn("Todas las empresas (consolidado)", textos(wb["CARÁTULA"]))
        ws = wb["LIBRO DIARIO"]
        fila, _ = buscar(ws, "TOTAL", col=1)
        self.assertEqual(ws.cell(row=fila, column=9).value, 0)
        self.assertTrue(sello["totales_control"]["cuadra"])

    def test_relacion_de_transacciones(self):
        d = datos_periodo()
        datos = {"empresas": [], "alias": d["alias"], "txs": d["txs"][:2],
                 "asientos": [g for g in d["asientos"] if g["tx_id"] in (1, 2)], "faltantes": [99], "truncado": False}
        r = ex.normalizar_receta({"modo": "transacciones", "tx_ids": [1, 2, 99], "nombre": "Gastos viaje Medellín"})
        wb, sello = libro(datos=datos, r=r)
        self.assertEqual(wb.sheetnames, list(ex.HOJAS_TRANSACCIONES.values()))
        sop = wb["SOPORTES"]
        fila, _ = buscar(sop, "⚠ sin soporte")
        self.assertEqual(sop[f"A{fila}"].value, 2)
        fila1, _ = buscar(sop, "Soporte 1")
        self.assertEqual(sop[f"E{fila1}"].hyperlink.target, URL)
        adv = " | ".join(sello["advertencias"])
        self.assertIn("#99", adv)
        self.assertIn("sin soporte", adv)
        self.assertTrue(sello["totales_control"]["cuadra"])
        self.assertIn("Gastos viaje Medellín", textos(wb["CARÁTULA"]))
        self.assertEqual(ex.nombre_archivo(r, datos)[:40], "FINSYS_RELACION_GASTOS-VIAJE-MEDELLIN_20")


# ══ Recolección (fuentes parcheadas) ══════════════════════════════════════

def _bp(cuentas):
    return {"cuentas": [{"cuenta_codigo": c, "cuenta_nombre": n, "cuenta_tipo": t, "saldo_inicial_db": idb,
                         "saldo_inicial_cr": icr, "mov_debito": d, "mov_credito": cr}
                        for c, n, t, idb, icr, d, cr in cuentas]}


class TestRecolectar(unittest.TestCase):

    def _parches(self, txs=None, portafolios=None, asientos=None):
        portafolios = portafolios if portafolios is not None else [
            {"id": 1, "name": "Negocio A"}, {"id": 2, "name": "Pegasus"}]
        bps = iter([_bp([("1110", "Bancos", "ACTIVO", 100, 0, 10, 0)]),
                    _bp([("1110", "Bancos", "ACTIVO", 0, 30, 0, 5)])])
        vacio = {"cuentas": []}
        p = [
            mock.patch("database_driver.obtener_portafolios", return_value=portafolios),
            mock.patch("database_driver.obtener_transacciones", return_value=txs or []),
            mock.patch("database_driver.listar_cartera", return_value=[]),
            mock.patch("kernel.kernel_reports.balance_prueba", side_effect=lambda *a: next(bps)),
            mock.patch("kernel.kernel_reports.estado_resultados", return_value={"ingresos": vacio, "gastos": vacio}),
            mock.patch("kernel.kernel_reports.balance_general",
                       return_value={"activos": vacio, "pasivos": vacio, "patrimonio": vacio, "utilidad_ejercicio": 0}),
            mock.patch("kernel.kernel_journal_workflow.obtener_asientos_agrupados",
                       return_value={"items": asientos or [], "total": len(asientos or [])}),
            mock.patch.object(ex, "_identidad_empresas", return_value={2: {"nombre": "Pegasus SAS", "nit": "901"}}),
            mock.patch.object(ex, "_nombres_oficiales", return_value={}),
        ]
        for x in p:
            x.start()
            self.addCleanup(x.stop)

    def test_varias_empresas_se_fusionan_con_saldo_neto(self):
        self._parches()
        d = ex.recolectar(receta(portfolio_id=None, portfolios=[1, 2]))
        self.assertEqual(d["bp"], [{"codigo": "1110", "nombre": "Bancos", "tipo": "ACTIVO",
                                    "ini": 70.0, "db": 10.0, "cr": 5.0}])
        self.assertEqual([e["nombre"] for e in d["empresas"]], ["Negocio A", "Pegasus SAS"])
        self.assertEqual(d["empresas"][1]["nit"], "901")

    def test_filtra_por_empresa_y_periodo_y_quita_duplicados(self):
        txs = [_tx(1, "INGRESO", 5, date(2026, 9, 3)), _tx(1, "INGRESO", 5, date(2026, 9, 3)),
               _tx(2, "GASTO", 5, date(2026, 8, 31)), _tx(3, "GASTO", 5, date(2026, 9, 9), portfolio_name="Pegasus")]
        self._parches(txs=txs)
        d = ex.recolectar(receta())
        self.assertEqual([t["id"] for t in d["txs"]], [1])

    def test_empresa_inexistente_es_error_del_usuario(self):
        self._parches()
        with self.assertRaises(ValueError):
            ex.recolectar(receta(portfolio_id=77))

    def test_sin_empresas_legibles_es_error_de_bd(self):
        self._parches(portafolios=[])
        with self.assertRaises(RuntimeError):
            ex.recolectar(receta())

    def test_transacciones_de_otra_empresa_se_rechazan(self):
        self._parches(txs=[_tx(1, "GASTO", 5, date(2026, 9, 3)), _tx(3, "GASTO", 5, date(2026, 9, 9), portfolio_name="Pegasus")])
        r = ex.normalizar_receta({"modo": "transacciones", "tx_ids": [1, 3], "portfolio_id": 1})
        with self.assertRaises(ValueError) as e:
            ex.recolectar(r)
        self.assertIn("#3", str(e.exception))

    def test_transacciones_reporta_faltantes_y_trae_solo_sus_asientos(self):
        asientos = [_grupo("JE-1", "CONTABILIZADO", 1, "2026-09-03", [("1110", "B", 5, 0), ("4135", "V", 0, 5)]),
                    _grupo("JE-9", "CONTABILIZADO", 9, "2026-09-03", [("1110", "B", 7, 0), ("4135", "V", 0, 7)])]
        self._parches(txs=[_tx(1, "INGRESO", 5, date(2026, 9, 3))], asientos=asientos)
        d = ex.recolectar(ex.normalizar_receta({"modo": "transacciones", "tx_ids": [1, 50]}))
        self.assertEqual(d["faltantes"], [50])
        self.assertEqual([g["entry_group_id"] for g in d["asientos"]], ["JE-1"])

    def test_nombres_el_plan_de_la_empresa_gana_al_puc(self):  # D-134-09
        self._parches()
        with mock.patch.object(ex, "_nombres_oficiales", return_value={"111005": "CUENTAS CORRIENTES"}):
            d = ex.recolectar(receta())
        self.assertEqual(d["nombres_cuenta"]["111005"], "CUENTAS CORRIENTES")
        self.assertEqual(d["nombres_cuenta"]["130505"], "Clientes Nacionales")   # del PUC estándar

    def test_ninguna_transaccion_existe(self):
        self._parches(txs=[])
        with self.assertRaises(ValueError):
            ex.recolectar(ex.normalizar_receta({"modo": "transacciones", "tx_ids": [1]}))


# ══ Endpoint ══════════════════════════════════════════════════════════════

class TestEndpoint(unittest.TestCase):  # CA-134-03

    @classmethod
    def setUpClass(cls):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from routers.analytics import router
        from routers.auth_guard import create_session_token
        app = FastAPI()
        app.include_router(router)
        cls.client = TestClient(app)

        def token(role):
            return {"Authorization": "Bearer " + create_session_token({"id": 1, "name": "Prueba", "role": role})}
        cls.contador, cls.member, cls.owner = token("contador"), token("member"), token("owner")

    def test_sin_sesion_401(self):
        self.assertEqual(self.client.get("/api/analytics/export.xlsx").status_code, 401)

    def test_rol_member_403(self):
        self.assertEqual(self.client.get("/api/analytics/export.xlsx", headers=self.member).status_code, 403)

    def test_contador_descarga_el_libro(self):
        sello = {"n_txs": 3, "n_asientos": 2, "totales_control": {"cuadra": True}, "advertencias": ["x"]}
        with mock.patch.object(ex, "armar_libro", return_value=(b"PK-xlsx", sello, "FINSYS_X.xlsx")) as armar:
            r = self.client.get("/api/analytics/export.xlsx?portfolio_id=1&desde=2026-09-01&hojas=diario&categorias=Ventas",
                                headers=self.contador)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers["content-type"], ex.MIME_XLSX)
        self.assertIn('filename="FINSYS_X.xlsx"', r.headers["content-disposition"])
        self.assertEqual(r.headers["x-finsys-cuadra"], "1")
        self.assertEqual(r.content, b"PK-xlsx")
        receta_enviada = armar.call_args.args[0]
        self.assertEqual(receta_enviada["portfolio_id"], 1)
        self.assertEqual(receta_enviada["filtros"]["categorias"], "Ventas")
        self.assertEqual(armar.call_args.kwargs["usuario"]["role"], "contador")

    def test_receta_invalida_400(self):
        r = self.client.get("/api/analytics/export.xlsx?hojas=inventada", headers=self.owner)
        self.assertEqual(r.status_code, 400)
        self.assertIn("Hojas desconocidas", r.json()["detail"])


# ══ 06-oct: comparativo, certificación (Ley 222 art. 37) y folio inicial ═══

def _con_comp(datos=None):
    d = datos or datos_periodo()
    d["comp"] = {
        "desde": date(2026, 8, 1), "hasta": date(2026, 8, 31),
        "er": {"ingresos": [{"codigo": "4135", "nombre": "Ventas", "tipo": "INGRESO", "saldo": 800_000.0},
                            {"codigo": "4210", "nombre": "Financieros", "tipo": "INGRESO", "saldo": 5_000.0}],
               "gastos": []},
        "bg": {"activos": [{"codigo": "1110", "nombre": "Bancos", "tipo": "ACTIVO", "saldo": 900_000.0}],
               "pasivos": [], "patrimonio": [{"codigo": "3105", "nombre": "Capital", "tipo": "PATRIMONIO", "saldo": 500_000.0}],
               "utilidad": 400_000.0},
    }
    return d


class TestComparativoYCertificacion(unittest.TestCase):

    def test_periodo_comparativo(self):
        pc = ex.periodo_comparativo
        self.assertEqual(pc(date(2026, 9, 1), date(2026, 9, 30), "periodo_anterior"), (date(2026, 8, 1), date(2026, 8, 31)))
        self.assertEqual(pc(date(2026, 1, 1), date(2026, 3, 31), "periodo_anterior"), (date(2025, 10, 1), date(2025, 12, 31)))
        self.assertEqual(pc(date(2026, 9, 10), date(2026, 9, 19), "periodo_anterior"), (date(2026, 8, 31), date(2026, 9, 9)))
        self.assertEqual(pc(date(2024, 2, 1), date(2024, 2, 29), "anio_anterior"), (date(2023, 2, 1), date(2023, 2, 28)))
        self.assertEqual(pc(date(2026, 1, 1), date(2026, 12, 31), "anio_anterior"), (date(2025, 1, 1), date(2025, 12, 31)))
        self.assertEqual(pc(None, date(2026, 9, 30), "periodo_anterior"), (None, date(2025, 9, 30)))

    def test_receta(self):
        r = receta(comparativo="anio_anterior", certificacion={"contador": " Ana ", "otra": "x"}, folio_inicial="41")
        self.assertEqual((r["comparativo"], r["folio_inicial"]), ("anio_anterior", 41))
        self.assertEqual(r["certificacion"], {"representante": "", "documento_representante": "", "contador": "Ana",
                                              "tarjeta_profesional": ""})
        self.assertEqual(receta(certificacion=True)["certificacion"]["contador"], "")
        self.assertIsNone(receta()["comparativo"])
        for malo in ({"comparativo": "siempre"}, {"folio_inicial": "x"}, {"folio_inicial": 0}):
            with self.assertRaises(ValueError):
                receta(**malo)
        with self.assertRaises(ValueError):
            ex.normalizar_receta({"modo": "transacciones", "tx_ids": [1], "comparativo": "anio_anterior"})

    def test_estados_comparativos(self):
        wb, sello = libro(_con_comp(), receta(comparativo="periodo_anterior"))
        er = wb["ESTADO DE RESULTADOS"]
        self.assertEqual(er["A1"].value, "ESTADO DE RESULTADOS COMPARATIVO")
        self.assertEqual([er.cell(row=4, column=j).value for j in (1, 2, 5, 6)], ["Código", "Cuenta", "Variación", "Var. %"])
        filas = {er.cell(row=i, column=1).value: i for i in range(1, er.max_row + 1)}
        i = filas["4135"]
        self.assertEqual((er[f"C{i}"].value, er[f"D{i}"].value, er[f"E{i}"].value), (1_000_000, 800_000, f"=C{i}-D{i}"))
        self.assertEqual(er[f"D{filas['4210']}"].value, 5_000)             # cuenta que solo existía antes
        self.assertEqual(er[f"C{filas['4210']}"].value, 0)
        bg = wb["BALANCE GENERAL"]
        textos = [str(c.value) for fila in bg.iter_rows() for c in fila if c.value]
        self.assertIn("BALANCE GENERAL COMPARATIVO", textos)
        self.assertTrue(any("Al 31/08/2026" in t for t in textos))
        self.assertEqual(sello["comparativo"], {"modo": "periodo_anterior", "desde": "2026-08-01", "hasta": "2026-08-31"})
        caratula = [str(c.value) for fila in wb["CARÁTULA"].iter_rows() for c in fila if c.value]
        self.assertTrue(any(t.startswith("período anterior: 01/08/2026") for t in caratula))

    def test_sin_comparativo_no_cambia(self):
        wb, sello = libro()
        self.assertEqual(wb["ESTADO DE RESULTADOS"]["A1"].value, "ESTADO DE RESULTADOS")
        self.assertIsNone(sello["comparativo"])
        self.assertNotIn("CERTIFICACIÓN", wb.sheetnames)

    def test_certificacion_y_folio_inicial(self):
        wb, sello = libro(r=receta(certificacion={"representante": "Julián Díaz", "tarjeta_profesional": "123456-T"},
                                   folio_inicial=41))
        self.assertIn("CERTIFICACIÓN", wb.sheetnames)
        ws = wb["CERTIFICACIÓN"]
        textos = " ".join(str(c.value) for fila in ws.iter_rows() for c in fila if c.value)
        self.assertIn("Artículo 37 de la Ley 222 de 1995", textos)
        self.assertIn("se han tomado fielmente de los libros", textos)
        self.assertIn("Finanzas Julian, NIT 900.123.456-7", textos)
        self.assertIn("Nombre: Julián Díaz", textos)
        self.assertIn("T.P.: 123456-T", textos)
        self.assertTrue(sello["certificacion"])
        self.assertEqual(sello["folio_inicial"], 41)
        self.assertEqual(wb["LIBRO DIARIO"].oddFooter.right.text, "Folio &P+40")
        self.assertEqual(wb["MAYOR"].oddFooter.right.text, "Folio &P+40")
        self.assertEqual(libro(r=receta(folio_inicial=1))[0]["LIBRO DIARIO"].oddFooter.right.text, "Folio &P")


class TestPatrimonioYFlujos(unittest.TestCase):   # 06-oct (NIIF Pymes §6 y §7)

    def test_clasificar_flujos_cuadra_por_partida_doble(self):
        bp = [
            {"codigo": "1110", "ini": 500.0, "db": 1000.0, "cr": 650.0},     # efectivo: +350
            {"codigo": "1305", "ini": 0.0, "db": 200.0, "cr": 0.0},          # deudores suben → −200
            {"codigo": "1524", "ini": 0.0, "db": 400.0, "cr": 0.0},          # compra de equipo → −400 (inversión)
            {"codigo": "1592", "ini": 0.0, "db": 0.0, "cr": 50.0},           # depreciación → +50 (no monetaria)
            {"codigo": "2105", "ini": 0.0, "db": 0.0, "cr": 300.0},          # préstamo → +300 (financiación)
            {"codigo": "2335", "ini": 0.0, "db": 0.0, "cr": 100.0},          # cuentas por pagar → +100
            {"codigo": "4135", "ini": 0.0, "db": 0.0, "cr": 1000.0},
            {"codigo": "5105", "ini": 0.0, "db": 450.0, "cr": 0.0},
            {"codigo": "5160", "ini": 0.0, "db": 50.0, "cr": 0.0},
            {"codigo": "ZZ", "ini": 0.0, "db": 0.0, "cr": 0.0},
        ]
        fl = ex.clasificar_flujos(bp)
        s = fl["secciones"]
        self.assertEqual(fl["utilidad"], 500.0)
        self.assertEqual(s["inv"], [("Propiedades, planta y equipo (15)", -400.0)])
        self.assertEqual(s["fin"], [("Obligaciones financieras (21)", 300.0)])
        self.assertEqual(dict(s["op_ct"]), {"Deudores (13)": -200.0, "Cuentas por pagar (23)": 100.0})
        self.assertEqual(s["op_aj"][0][1], 50.0)
        neto = fl["utilidad"] + sum(v for k in ("op_aj", "op_ct", "inv", "fin", "otros") for _e, v in s[k])
        self.assertEqual((fl["efectivo_ini"], fl["efectivo_fin"]), (500.0, 850.0))
        self.assertAlmostEqual(neto, fl["efectivo_fin"] - fl["efectivo_ini"])
        self.assertEqual(fl["sin_clasificar"], [])

    def test_hojas_en_el_libro(self):
        wb, sello = libro(r=receta(hojas=["estado_resultados", "balance_general", "cambios_patrimonio", "flujos_efectivo"],
                                   certificacion=True))
        cp = wb["CAMBIOS EN EL PATRIMONIO"]
        filas = {cp.cell(row=i, column=2).value: i for i in range(1, cp.max_row + 1)}
        i = filas["Resultado del ejercicio (sin asiento de cierre)"]
        self.assertEqual((cp[f"C{i}"].value, cp[f"D{i}"].value, cp[f"E{i}"].value), (0, 800_000, 0))
        self.assertEqual(cp[f"F{i}"].value, f"=C{i}+D{i}-E{i}")
        textos_cp = [str(c.value) for f in cp.iter_rows() for c in f if c.value]
        self.assertIn("Patrimonio + utilidad según el BALANCE GENERAL", textos_cp)
        fe = wb["FLUJOS DE EFECTIVO"]
        textos = {fe.cell(row=k, column=1).value: fe.cell(row=k, column=2).value for k in range(1, fe.max_row + 1)}
        self.assertEqual(textos["Utilidad (pérdida) del período"], 800_000)
        self.assertEqual(textos["Efectivo y equivalentes al inicio del período (grupo 11)"], 500_000)
        self.assertEqual(textos["Efectivo al final según el balance de prueba (grupo 11)"], 1_300_000)
        cert = " ".join(str(c.value) for f in wb["CERTIFICACIÓN"].iter_rows() for c in f if c.value)
        self.assertIn("el estado de flujos de efectivo", cert)
        self.assertIn("FLUJOS DE EFECTIVO", sello["hojas"])


if __name__ == "__main__":
    unittest.main()
