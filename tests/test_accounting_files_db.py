# -*- coding: utf-8 -*-
"""Integración del organizador contable 13.5 contra Supabase (CA-135-02/03/05/08).

ESCRIBE en accounting_files (tabla nueva que nadie más lee) y borra lo suyo al
terminar. Por eso es OPT-IN: solo corre con FINSYS_TEST_BD_ESCRIBE=1 y BD a mano.
Cada corrida consume folios de la secuencia (un folio no se reutiliza jamás).

  · generar → descargar: el SHA-256 de la descarga = el del archivo generado (CA-135-03)
  · la carátula del archivo lleva el folio; la ficha recién generada está ✅ VIGENTE
  · pre-vuelo = sello del archivo generado (CA-135-02)
  · regenerar crea folio nuevo con reemplaza_a y el original queda intacto (CA-135-05)
  · relación de 3 TXs: exactamente esas 3 y sus asientos cuadran (CA-135-08, parte backend)

Ejecutar:  set FINSYS_TEST_BD_ESCRIBE=1 & .venv\\Scripts\\python.exe -m unittest tests.test_accounting_files_db -v
"""
import hashlib
import io
import os
import sys
import unittest

from openpyxl import load_workbook

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "fin_sys_core"))

_env = os.path.join(_ROOT, ".env")
if os.path.exists(_env):
    with open(_env, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, _, v = line.partition("=")
                os.environ.setdefault(k.strip(), v.strip())

import fin_sys_core  # noqa: E402,F401
from fin_sys_core import accounting_files_driver as drv  # noqa: E402

USUARIO = {"name": "test-organizador", "uid": "0", "role": "owner"}
COMPARAR = ("n_txs", "n_asientos", "n_lineas", "totales_control", "advertencias", "hojas", "empresas")


def _bd_lista():
    if os.getenv("FINSYS_TEST_BD_ESCRIBE") != "1":
        return False
    try:
        from fin_sys_core.db_pool import get_conn, put_conn
        conn = get_conn()
        try:
            cur = conn.cursor()
            cur.execute("SELECT 1 FROM accounting_files LIMIT 1")
            cur.close()
            return True
        finally:
            put_conn(conn)
    except Exception:
        return False


def _consulta(sql, params=()):
    from fin_sys_core.db_pool import get_conn, put_conn
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute(sql, params)
        filas = cur.fetchall()
        cur.close()
        conn.rollback()
        return filas
    finally:
        put_conn(conn)


@unittest.skipUnless(_bd_lista(), "opt-in: FINSYS_TEST_BD_ESCRIBE=1 y migración 13.5 aplicada")
class TestOrganizadorBD(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.creados = []
        filas = _consulta("""
            SELECT portfolio_id, date_trunc('month', MAX(transaction_date))::date
              FROM transactions GROUP BY portfolio_id ORDER BY COUNT(*) DESC LIMIT 1
        """)
        cls.pid, inicio = filas[0]
        import calendar
        fin = inicio.replace(day=calendar.monthrange(inicio.year, inicio.month)[1])
        cls.receta = {"portfolio_id": cls.pid, "desde": inicio.isoformat(), "hasta": fin.isoformat()}
        cls.tx_ids = [r[0] for r in _consulta(
            "SELECT id FROM transactions WHERE portfolio_id = %s ORDER BY id DESC LIMIT 3", (cls.pid,))]

    @classmethod
    def tearDownClass(cls):
        for i in reversed(cls.creados):
            try:
                drv.eliminar(i)
            except drv.NoEncontrado:
                pass

    def _generar(self, **kw):
        out = drv.generar(kw.pop("receta", self.receta), USUARIO, **kw)
        self.creados.append(out["id"])
        return out

    def test_generar_descargar_regenerar(self):
        g = self._generar(paquete="cierre_mes")
        self.assertRegex(g["folio"], r"^EXP-\d{4}-\d{4,}$")

        d = drv.descargar(g["id"])
        self.assertEqual(hashlib.sha256(d["contenido"]).hexdigest(), g["sha256"])     # CA-135-03
        self.assertEqual(d["sha256"], g["sha256"])
        self.assertEqual(drv.descargar(g["id"])["contenido"], d["contenido"])
        wb = load_workbook(io.BytesIO(d["contenido"]), read_only=True)
        celdas = [c for fila in wb[wb.sheetnames[0]].iter_rows(values_only=True) for c in fila if c]
        self.assertIn(g["folio"], celdas)
        self.assertEqual(wb.sheetnames[1:], ["LIBRO DIARIO", "MAYOR", "BALANCE DE PRUEBA",
                                             "ESTADO DE RESULTADOS", "BALANCE GENERAL"])
        wb.close()

        f = drv.ficha(g["id"])
        self.assertEqual(f["vigencia"]["estado"], "VIGENTE")
        self.assertEqual(f["descargas"], 2)
        self.assertEqual(f["paquete"], "cierre_mes")
        self.assertEqual(f["portfolio_id"], self.pid)

        p = drv.prevuelo(self.receta, "cierre_mes", usuario=USUARIO)                   # CA-135-02
        for k in COMPARAR:
            self.assertEqual(p["sello"][k], g["sello"][k], k)

        n = drv.regenerar(g["id"], USUARIO)                                            # CA-135-05
        self.creados.append(n["id"])
        self.assertNotEqual(n["folio"], g["folio"])
        self.assertEqual(drv.ficha(n["id"])["reemplaza"], {"id": g["id"], "folio": g["folio"]})
        original = drv.ficha(g["id"])
        self.assertEqual(original["sha256"], g["sha256"])
        self.assertEqual([x["id"] for x in original["reemplazado_por"]], [n["id"]])

        lista = drv.listar({"q": g["folio"]})
        self.assertEqual([i["id"] for i in lista["items"]], [g["id"]])
        self.assertGreaterEqual(drv.resumen()["archivos"], 2)

    def test_relacion_de_tres_transacciones(self):                                     # CA-135-08
        r = self._generar(receta={"modo": "transacciones", "tx_ids": self.tx_ids, "nombre": "test-relacion"})
        self.assertEqual(r["paquete"], "relacion")
        self.assertEqual(r["sello"]["n_txs"], len(self.tx_ids))
        self.assertTrue(r["sello"]["totales_control"]["cuadra"])
        f = drv.ficha(r["id"])
        self.assertEqual(f["vigencia"]["estado"], "VIGENTE")
        self.assertEqual(sorted(f["receta"]["tx_ids"]), sorted(self.tx_ids))
        v = drv.vista_previa(r["id"], "RELACIÓN", 10)
        self.assertEqual(v["hoja"], "RELACIÓN")
        self.assertIn("RELACIÓN", v["hojas"])


if __name__ == "__main__":
    unittest.main()
